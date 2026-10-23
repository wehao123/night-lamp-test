from pathlib import Path
from docx import Document
root=Path(__file__).resolve().parent
src=root/'室内分项实验综合报告_420–720nm修订版_v2_公式修订.docx'
out=root/'室内分项实验综合报告_420–720nm修订版_v3.docx'
doc=Document(src)
for p in doc.paragraphs:
    if p.text in ('4.3.2 非线性比值与相对偏差','B.3.2 非线性比值与相对偏差'):
        old=p.text
        for r in list(p.runs): r._element.getparent().remove(r._element)
        p.add_run(old.replace('非线性比值','比例响应比'))
    elif p.text == '定义非线性比值：':
        for r in list(p.runs): r._element.getparent().remove(r._element)
        p.add_run('定义比例响应比（本报告指标）：')
doc.save(out)
print(out)
