from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE = Path(r"C:\Users\haowei\Desktop\夜间水质资料\夜间水质资料\光谱仪室内实验")
FIG = BASE / "figures"
OUT = Path(r"C:\Users\haowei\Documents\ChatGPT\灯光距离实验\OCR-25F02100_室内实验B_C_D完整报告.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def keep_together(paragraph):
    paragraph.paragraph_format.keep_together = True


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.25)
sec.bottom_margin = Cm(2.1)
sec.left_margin = Cm(2.35)
sec.right_margin = Cm(2.15)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "宋体"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.35
normal.paragraph_format.space_after = Pt(5)

for name, size, font in [("Title", 22, "黑体"), ("Heading 1", 16, "黑体"), ("Heading 2", 13, "黑体"), ("Heading 3", 11.5, "黑体")]:
    st = styles[name]
    st.font.name = font
    st._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(6)

styles["Title"].paragraph_format.space_after = Pt(12)
styles["Caption"].font.name = "宋体"
styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Caption"].font.size = Pt(9)
styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)


def p(text="", bold_lead=None, indent=True):
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para


def bullet(text):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.left_indent = Cm(0.75)
    para.paragraph_format.first_line_indent = Cm(-0.25)
    para.add_run(text)
    return para


def numbered(text):
    para = doc.add_paragraph(style="List Number")
    para.paragraph_format.left_indent = Cm(0.75)
    para.paragraph_format.first_line_indent = Cm(-0.25)
    para.add_run(text)
    return para


def equation(text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(5)
    run = para.add_run(text)
    run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
    run.font.size = Pt(11)
    keep_together(para)
    return para


def table(headers, rows, widths=None, font_size=8.5):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.style = "Table Grid"
    hdr = tbl.rows[0]
    repeat_table_header(hdr)
    for j, h in enumerate(headers):
        cell = hdr.cells[j]
        cell.text = str(h)
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(font_size)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell)
    for i, row in enumerate(rows):
        cells = tbl.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2:
                set_cell_shading(cells[j], "EEF4F8")
            for para in cells[j].paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 or len(headers) <= 3 else WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(font_size)
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            set_cell_margins(cells[j])
    if widths:
        for row in tbl.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return tbl


def figure(filename, caption, width_cm=16.0):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG / filename), width=Cm(width_cm))
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_together = True
    cap.paragraph_format.space_after = Pt(8)
    cap.add_run(caption)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# 封面
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("二代 OCR 模组室内实验分析报告")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
r = subtitle.add_run("实验 B 积分时间线性度  实验 C 信噪比  实验 D 弱光性能")
r.font.name = "黑体"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体"); r.font.size = Pt(13)
table(["项目", "内容"], [
    ["探头型号", "OCR-25F02100-1"],
    ["标定 ID", "1"],
    ["采集日期", "2026-09-20"],
    ["报告日期", "2026-09-23"],
    ["数据范围", "32 份 CSV；288 个通道；305.883–878.092 nm"],
], widths=[4.0, 11.5], font_size=10)
p("本报告汇总积分时间相对线性度、信噪比和弱光性能三项室内实验结果。全部定量结果均由原始重复帧计算，未进行平滑、滤波或插值。", indent=False)
page_break()

# B
doc.add_heading("实验 B  积分时间线性度", level=1)
doc.add_heading("B.1 实验目的", level=2)
p("评价光谱仪在光源与光路条件保持不变时，净光响应随积分时间变化的比例关系，分析不同码值区间内的响应偏差，为积分时间选择和工作码值范围提供依据。")
p("本次采用长、短积分时间配对比较的方法。该方法评价已测条件下的相对响应一致性，不等同于覆盖完整量程的多点线性标定。")

doc.add_heading("B.2 实验数据与预处理", level=2)
p("本次使用12份N模式数据，组成6组长、短积分时间配对。每份数据包含288个原生通道，波长范围为305.883–878.092 nm。")
table(["配对组", "短积分时间", "长积分时间", "实际时间倍率", "短档/长档帧数"], [
    ["3W，最大64 ms组", "32 ms", "64 ms", "2.0000", "30/30"],
    ["4W，最大64 ms组", "32 ms", "64 ms", "2.0000", "30/30"],
    ["5W，最大64 ms组", "32 ms", "64 ms", "2.0000", "50/50"],
    ["3W，最大128 ms组", "64 ms", "128 ms", "2.0000", "30/30"],
    ["4W，最大128 ms组", "64 ms", "128 ms", "2.0000", "30/30"],
    ["5W，最大128 ms组", "62 ms", "128 ms", "2.0645", "30/30"],
], widths=[4.2, 2.7, 2.7, 2.8, 3.1])
p("“3W、4W、5W”作为原始分组标签保留，不作为实测光功率参与计算。文件名末尾的64或128表示该组最大积分时间。")
p("其中，5W、最大128 ms组的短档文件名标注为64 ms，但文件内部记录的实际积分时间为62 ms。本报告采用实际记录值，计算倍率为128/62，而非2。")
p("N模式数据已经由设备完成扣暗处理，导出时额外包含1000 DN固定偏置。因此，每帧净光响应按下式计算：")
equation("Sᵢ(λ,t) = Nᵢ(λ,t) − 1000")
p("其中，Nᵢ为导出的N模式码值，Sᵢ为扣除固定偏置后的净光响应。N模式数据不再减去独立暗场，以免重复扣暗。全部统计使用原始重复帧，不进行平滑、滤波或插值。")

doc.add_heading("B.3 评价方法", level=2)
doc.add_heading("B.3.1 积分时间比例关系", level=3)
p("在入射光保持稳定、仪器响应理想线性的条件下，同一波长的净光响应应与积分时间成正比：")
equation("S(λ,t) = K(λ)t")
p("其中，K(λ)综合反映当前入射光与仪器响应。分别计算长、短积分时间下的平均净光响应，并以短积分实测均值为参考，推算长积分时间下的理论响应：")
equation("S̄短(λ) = (1/n短) Σ Sᵢ(λ,t短)")
equation("S̄长(λ) = (1/n长) Σ Sᵢ(λ,t长)")
equation("S理论,长(λ) = S̄短(λ) · t长/t短")
p("此处“理论响应”是根据短积分测量结果和时间倍率得到的比例外推值，并非独立标准仪器给出的绝对真值。")

doc.add_heading("B.3.2 非线性比值与相对偏差", level=3)
p("定义非线性比值：")
equation("γ(λ) = S理论,长(λ) / S̄长(λ)")
bullet("γ=1：实测响应与积分时间比例关系一致；")
bullet("γ>1：长积分实测响应低于理论外推值；")
bullet("γ<1：长积分实测响应高于理论外推值。")
p("为便于以百分比表达偏差，定义：")
equation("δ(λ) = [S̄长(λ) − S理论,长(λ)] / S理论,长(λ) × 100%")
equation("δ = (1/γ − 1) × 100%")
p("δ>0表示实测高于理论值，δ<0表示实测低于理论值。由于分母不同，|δ|与|γ−1|×100%不完全相等，报告中分别标明。")

doc.add_heading("B.3.3 统计通道筛选", level=3)
p("弱信号下，噪声容易放大长、短积分响应比值的波动。纳入线性度统计的通道需同时满足：")
equation("S̄短 > 0，S̄长 > 0，SNR短 ≥ 20，SNR长 ≥ 20")
p("各档SNR由净光均值除以帧间样本标准差计算。未满足上述条件的通道不纳入本章汇总统计，但不据此认定其存在器件非线性。")
table(["指标", "含义"], [
    ["δ中位数", "描述整体偏差方向与典型水平"],
    ["|δ|的95分位", "描述大多数入选通道的绝对偏差水平"],
    ["最大|δ|", "描述入选通道中观察到的最大相对偏差"],
    ["最大|γ−1|", "描述非线性比值偏离1的最大程度"],
], widths=[4.3, 11.2], font_size=9.5)
p("95分位是样本分布统计量，不是95%置信区间。本章的δ是相对比例外推值的偏差，不是多点直线拟合残差。")

doc.add_heading("B.4 实验结果", level=2)
doc.add_heading("B.4.1 各配对组结果", level=3)
table(["配对组", "入选通道数", "δ中位数", "|δ|95分位", "最大|δ|", "最大|γ−1|"], [
    ["32→64 ms，3W", "256/288", "+0.466%", "1.528%", "5.014%", "5.278%"],
    ["32→64 ms，4W", "286/288", "+0.071%", "1.810%", "6.288%", "6.710%"],
    ["32→64 ms，5W", "288/288", "−0.155%", "2.214%", "4.460%", "4.668%"],
    ["64→128 ms，3W", "254/288", "+0.562%", "1.859%", "4.975%", "5.235%"],
    ["64→128 ms，4W", "272/288", "+0.402%", "1.865%", "5.141%", "4.890%"],
    ["62→128 ms，5W", "285/288", "−0.144%", "2.134%", "4.082%", "3.921%"],
], widths=[3.8, 2.5, 2.2, 2.3, 2.3, 2.4], font_size=8.2)
p("六组的δ中位数介于−0.155%至+0.562%，表明各组典型响应接近积分时间比例关系。但部分通道仍存在较大偏差，最大绝对相对偏差为4.082%–6.288%，因此不能仅凭中位数接近零判断所有通道均满足同一精度要求。")

doc.add_heading("B.4.2 码值分段结果", level=3)
p("以长积分实测平均净光响应为分段依据，汇总六组中通过筛选的“通道×组”数据点。")
table(["长积分净光区间 / DN", "数据点数", "涉及组数", "δ中位数", "|δ|95分位", "最大|δ|"], [
    ["[0, 4000)", "292", "6", "−0.187%", "3.590%", "6.288%"],
    ["[4000, 8000)", "81", "6", "+0.475%", "1.192%", "1.738%"],
    ["[8000, 15000)", "204", "6", "+0.691%", "1.098%", "1.454%"],
    ["[15000, 25000)", "346", "6", "+0.751%", "1.149%", "1.420%"],
    ["[25000, 35000)", "347", "6", "+0.226%", "0.557%", "0.753%"],
    ["[35000, 45000)", "230", "4", "−0.461%", "1.113%", "1.504%"],
    ["[45000, 55000)", "141", "2", "−1.927%", "2.272%", "2.396%"],
], widths=[3.7, 2.1, 1.8, 2.4, 2.5, 2.4], font_size=8.2)
p("在4000–45000 DN区间，共有1208个入选数据点：")
equation("max|δ| = 1.738%，max|γ−1| × 100% = 1.708%")
p("上述结果表明，本次已测4000–45000 DN范围内的入选数据点，其相对比例偏差均小于2%。这里的2%用于描述实测结果；在未明确线性度验收标准前，不据此直接作出整机“线性度合格”的判断。")
p("25000–35000 DN区间的统计表现最佳：δ中位数为+0.226%，|δ|95分位为0.557%，最大绝对偏差为0.753%。该区间可作为后续参数选择的参考，但仍需同时检查亮场量程、目标波段SNR及漂移。这里的DN均指扣除固定偏置后的净光响应，不能直接替代实验C中用于检查80%量程的亮场码值。")

doc.add_heading("B.4.3 低码值与高码值区表现", level=3)
p("低于4000 DN时，响应比值的离散程度增大。该区间|δ|95分位为3.590%，最大值为6.288%。弱信号下，随机噪声及两次采集之间的光源变化对比值影响更明显，不能将全部偏差归因于器件非线性。", bold_lead="低于4000 DN时，响应比值的离散程度增大。")
p("45000 DN以上出现实测响应低于理论外推值的趋势。在45000–55000 DN区间，δ中位数为−1.927%，对应γ>1，与响应压缩表现相符。但该区间仅由两组数据覆盖，且本次未测得完整饱和过程，因此不报告精确饱和起点，也不将观察到的趋势直接等同于数字饱和。", bold_lead="45000 DN以上出现实测响应低于理论外推值的趋势。")

doc.add_heading("B.4.4 不同波段的分布", level=3)
p("将入选通道按约306–400、400–480、480–560及560–878 nm分组，在同一“净光码值—非线性比值”图中显示。不同波段覆盖的码值范围不一致：弱响应波段主要分布于低码值区，高响应波段延伸至较高码值区。因此，波长与码值的影响在现有数据中未完全分离。")
p("该图可用于展示各波段数据的分布，但不足以证明非线性仅由码值决定，或与波长无关。若需识别独立的波长影响，应在相近码值条件下对不同波段进行受控比较。")

doc.add_heading("B.5 图示", level=2)
figure("01_linearity_scatter.png", "图1  六组积分时间配对的非线性比值与净光码值关系。横轴为长积分实测平均净光响应，纵轴为理论响应与实测响应之比γ，虚线表示γ=1。")
figure("02_delta_by_bin.png", "图2  不同净光码值区间的相对偏差统计。蓝色柱形对应左轴，橙色折线对应右轴；95分位不代表最大偏差或置信区间。")
figure("03_band_split.png", "图3  不同波段的非线性比值与净光码值分布。由于各波段覆盖的码值范围不同，该图用于描述数据分布，不作为非线性与波长无关的证明。")

doc.add_heading("B.6 结论与适用范围", level=2)
numbered("已测中间码值区的响应接近积分时间比例关系。在长积分净光响应4000–45000 DN范围内，1208个入选数据点的最大绝对相对偏差为1.738%，最大非线性比值偏离为1.708%。")
numbered("25000–35000 DN区间在本次分段统计中表现最佳。该区间最大绝对相对偏差为0.753%，可作为后续工作码值选择的参考，但不保证所有波长及未测条件均具有相同表现。")
numbered("弱信号区的比值离散较大，高码值区出现响应压缩趋势。低于4000 DN的数据需结合SNR谨慎解释；45000 DN以上的数据覆盖有限，不能据此确定精确饱和起点。")
numbered("本次结果为积分时间配对的相对线性度评价。长、短积分采集并非同步，且缺少同步辐亮度记录，测得偏差可能同时包含仪器非线性与光源变化，尚不能分离两者贡献。")
numbered("本次不报告多点拟合的R²与最大相对残差。每组仅包含两个积分时间水平，对两个均值进行带截距直线拟合不能有效检验线性程度。若需完整标定，应在稳定光源下增加积分时间采样点，并覆盖弱信号、中间响应及接近饱和区域。")

page_break()

# C
doc.add_heading("实验 C  信噪比", level=1)
doc.add_heading("C.1 实验目的", level=2)
p("评价光谱仪在不同光源电压和积分时间条件下的逐波长信噪比、重复性及采集期漂移，并结合码值余量确定工作候选条件与可用波段。")
p("本实验光谱仪无可调增益，因此仅比较光源电压与积分时间，不设置增益变量。")

doc.add_heading("C.2 实验数据与预处理", level=2)
p("本次共采集5档光源电压、10组测量条件，每组包含同积分时间下的亮场和暗场数据。光谱覆盖288个原生通道，波长范围为305.883–878.092 nm。")
table(["光源电压", "积分时间", "每个条件亮场/暗场帧数"], [
    ["60 V", "64、128 ms", "30/30"], ["63 V", "64、128 ms", "30/30"],
    ["68 V", "32、64 ms", "30/30"], ["72 V", "32、64 ms", "30/30"],
    ["76 V", "32、64 ms", "32 ms：50/50；64 ms：30/30"],
], widths=[3.5, 4.5, 7.5], font_size=9.2)
p("光源通过调节电压减弱，但未记录参考板辐亮度。因此，电压仅作为实验条件标签，不等同于亮度，也不能将电压比视为辐亮度比。")
p("亮场导出值记为Lᵢ(λ)，暗场导出值记为Dⱼ(λ)。依据设备侧确认，两者均包含固定的1000 DN偏置。对每个测量条件，先计算平均暗谱，再对每帧亮场进行扣暗：")
equation("D̄(λ) = (1/nD) Σ Dⱼ(λ)")
equation("Sᵢ(λ) = Lᵢ(λ) − D̄(λ)")
p("由于亮场与暗场包含相同的固定偏置，相减后1000 DN自动抵消，净信号不再额外减去1000。全部统计沿重复帧方向计算，不沿波长方向计算，未进行平滑、滤波或插值。")

doc.add_heading("C.3 评价方法与判据", level=2)
doc.add_heading("C.3.1 信噪比与变异系数", level=3)
equation("S̄(λ) = (1/nL) Σ Sᵢ(λ)")
equation("sS(λ) = √{Σ[Sᵢ(λ)−S̄(λ)]²/(nL−1)}")
equation("SNR(λ) = S̄(λ)/sS(λ)")
equation("CV(λ) = sS(λ)/|S̄(λ)| × 100%")
p("SNR越高，表示信号相对于帧间波动越强；CV越低，表示重复测量的相对离散程度越小。对于正信号，CV=100%/SNR。因此，SNR≥20与CV≤5%等价，SNR≥50与CV≤2%等价，两者不是相互独立的性能证据。")
p("本章报告的是以帧间标准差为噪声估计的SNR，不乘以√nL换算成多帧均值的SNR。减去同一条平均暗谱不改变亮值帧间标准差；该统计反映本次序列的重复性，不包含平均暗谱估计误差带来的全部测量不确定度。")

doc.add_heading("C.3.2 码值余量", level=3)
p("为保留量程余量，各通道所有亮帧中的最大值扣除固定偏置后，应满足：")
equation("maxᵢ[Lᵢ(λ)] − 1000 ≤ 0.8 × 65535 = 52428 DN")
p("等价地，含偏置的亮场导出值应不超过53428 DN。此处使用亮场值检查量程占用，而非扣暗后的净信号。超过80%量程属于码值余量不足，不等同于已经发生数字饱和。")

doc.add_heading("C.3.3 采集期漂移", level=3)
p("根据实际时间戳，对各通道净信号进行线性拟合：")
equation("Sᵢ(λ) = a(λ) + k(λ)τᵢ + εᵢ")
equation("D_fit(λ) = k(λ)[τ_end − τ_start] / |S̄(λ)| × 100%")
p("要求|Dfit(λ)|≤2%。该指标描述采集期间的拟合变化，不等同于最大值与最小值之间的峰峰变化。弱信号下随机波动也可能导致拟合漂移超限，不能仅据此认定灯具发生持续漂移。")

doc.add_heading("C.3.4 联合筛选规则", level=3)
table(["层级", "筛选要求"], [
    ["工作候选通道", "净信号均值为正、无数字饱和、SNR≥20、CV≤5%、满足80%量程要求、拟合漂移绝对值≤2%"],
    ["推荐候选通道", "在工作候选要求基础上，进一步满足SNR≥50、CV≤2%"],
], widths=[3.6, 11.9], font_size=9.2)
p("未通过任一工作候选要求的通道，标记为该条件下的非工作候选通道，并保留具体原因。该标记针对当前测量条件，不表示该波长永久不可用。")

doc.add_heading("C.4 实验结果", level=2)
doc.add_heading("C.4.1 各条件信噪比与联合筛选结果", level=3)
table(["条件", "SNR中位", "SNR≥20", "SNR≥50", "超80%", "漂移超2%", "工作候选", "推荐候选"], [
    ["60 V / 64 ms", "98.3", "235", "207", "0", "49", "224", "205"],
    ["60 V / 128 ms", "140.4", "264", "223", "0", "39", "241", "223"],
    ["63 V / 64 ms", "111.0", "247", "217", "0", "37", "236", "217"],
    ["63 V / 128 ms", "163.4", "285", "230", "0", "24", "264", "230"],
    ["68 V / 32 ms", "95.1", "238", "211", "0", "42", "230", "210"],
    ["68 V / 64 ms", "139.0", "275", "227", "0", "20", "260", "227"],
    ["72 V / 32 ms", "112.9", "260", "219", "0", "33", "242", "219"],
    ["72 V / 64 ms", "167.1", "288", "235", "0", "16", "272", "235"],
    ["76 V / 32 ms", "128.3", "271", "230", "0", "16", "264", "230"],
    ["76 V / 64 ms", "191.3", "288", "247", "78", "11", "199", "168"],
], widths=[3.0, 1.8, 1.7, 1.7, 1.5, 1.8, 1.8, 1.8], font_size=7.5)
p("各条件均包含288个通道。同一通道可能同时触发多项不通过条件，因此超限数量不能直接相加。")
p("10个条件的逐波长SNR中位数为95.1–191.3，但中位数不能代表每个通道均达标。60 V / 64 ms条件下，53个通道SNR低于20，主要集中于短波端；延长积分时间至128 ms后，不达标通道减少至24个。")
p("72 V / 64 ms条件下，全体288个通道SNR均≥20，且均满足80%量程要求。进一步考虑漂移后，272个通道通过工作候选筛选，235个通道通过推荐候选筛选，是本次已测条件中工作候选通道数最多的条件。")

doc.add_heading("C.4.2 逐波长分布", level=3)
p("逐波长SNR曲线显示，约500–800 nm中段的信噪比较高，短波端信噪比相对较低；CV曲线呈相反趋势。图中的曲线展示所有通道，尚未通过码值余量和漂移联合筛选。因此，曲线高于SNR阈值不代表对应通道已经满足全部工作要求。")
figure("04_snr_curves.png", "图4  不同电压与积分时间条件下的逐波长信噪比。纵轴采用对数刻度，水平线分别表示SNR=20和SNR=50；最终可用性由码值余量及漂移联合筛选确定。")
figure("05_cv_curves.png", "图5  不同电压与积分时间条件下的逐波长变异系数。纵轴采用对数刻度，水平线分别表示CV=5%和CV=2%；对于正信号，CV=100%/SNR。")

doc.add_heading("C.4.3 积分时间的影响", level=3)
equation("G(λ) = SNR长(λ) / SNR短(λ)")
table(["电压", "积分时间变化", "逐波长SNR增益中位数"], [
    ["60 V", "64→128 ms", "1.516"], ["63 V", "64→128 ms", "1.527"],
    ["68 V", "32→64 ms", "1.508"], ["72 V", "32→64 ms", "1.510"],
    ["76 V", "32→64 ms", "1.554"],
], widths=[3.2, 5.2, 7.1], font_size=9.2)
p("积分时间加倍后，逐波长SNR增益中位数为1.51–1.55，说明延长积分时间能够改善重复测量信噪比。但积分时间选择还需满足量程要求。")
p("76 V / 64 ms虽然具有最高的全谱SNR中位数，但589.2–740.2 nm范围内的78个通道超过80%量程，联合筛选后仅剩199个工作候选通道。同电压改用32 ms后，无通道超过80%量程，工作候选通道增加至264个。因此，不能仅以SNR最高或信号幅值最大作为参数选择依据。")

doc.add_heading("C.4.4 连续可用波段", level=3)
table(["条件", "最长连续工作候选波段 / nm", "通道数", "最长连续推荐候选波段 / nm", "通道数"], [
    ["60 V / 64 ms", "499.1–837.6", "181", "515.9–837.6", "174"],
    ["60 V / 128 ms", "474.6–878.1", "223", "474.6–878.1", "223"],
    ["63 V / 64 ms", "452.3–878.1", "232", "501.5–878.1", "212"],
    ["63 V / 128 ms", "447.2–878.1", "234", "464.7–878.1", "227"],
    ["68 V / 32 ms", "494.2–874.5", "212", "511.1–874.5", "205"],
    ["68 V / 64 ms", "437.1–878.1", "238", "464.7–878.1", "227"],
    ["72 V / 32 ms", "482.0–878.1", "220", "486.9–878.1", "218"],
    ["72 V / 64 ms", "409.0–878.1", "249", "447.2–878.1", "234"],
    ["76 V / 32 ms", "442.2–878.1", "236", "462.3–878.1", "228"],
    ["76 V / 64 ms", "741.9–878.1", "96", "741.9–878.1", "96"],
], widths=[2.7, 4.5, 1.6, 4.7, 1.6], font_size=7.5)
p("表中只列出每个条件下最长的一段连续候选波段，其余波段中仍可能存在通过筛选的通道。波段边界取原生通道中心波长，不表示边界已通过连续扫描确定。")

doc.add_heading("C.4.5 暗场与漂移表现", level=3)
p("各条件暗场均值的通道中位数为9174–9514 DN，扣除1000 DN固定偏置后为8174–8514 DN；暗场帧间标准差的通道中位数为31.5–35.0 DN。扣除固定偏置后的暗场响应仍可能包含电子学基线等成分，不能直接称为纯暗电流；暗场标准差也不能直接等同于纯读出噪声。")
p("以72 V / 64 ms为例，16个漂移超限通道全部位于约306–406 nm，净信号均值的中位数约为1269 DN，低于该条件的全谱中位数20648 DN。其相对漂移指标可能受到弱信号随机波动影响，现有结果不足以将超限原因单独归为光源漂移。")

doc.add_heading("C.5 结论与适用范围", level=2)
numbered("已测条件下信噪比总体良好，适当参数下可实现全波段SNR≥20。72 V / 64 ms条件下，全体288个通道同时满足基本SNR要求与80%量程要求。")
numbered("综合信噪比、码值余量与采集期漂移，72 V / 64 ms为本次测试的优选条件。272个通道通过工作候选筛选，最长连续工作候选波段为409.0–878.1 nm；最长连续推荐候选波段为447.2–878.1 nm。")
numbered("延长积分时间有助于提高SNR，但应同时检查量程余量。76 V / 64 ms虽具有最高SNR中位数，却存在78个通道超过80%量程，不能作为全波段优选设置。")
numbered("本次结论适用于现有室内数据。实际采集为30或50帧，持续约16.0–31.3 s，尚不能认定已完成“100帧内漂移≤2%”验证；因缺少辐亮度记录，也不能给出绝对亮度对应的参数表。室外实验可据此选择初始参数，最终工作波段、SNR及稳定性应由现场数据确认。")

page_break()

# D
doc.add_heading("实验 D  弱光性能", level=1)
doc.add_heading("D.1 实验目的", level=2)
p("评价光谱仪在光源逐级减弱时的逐波长净光响应、信噪比和重复性，分析延长积分时间对弱信号测量的改善效果，确定已测条件下的工作候选波段，并明确最低可测亮度及后续改进需求的验证范围。")
p("本实验光谱仪无可调增益，因此不设置增益变量。")

doc.add_heading("D.2 实验数据与条件", level=2)
p("本章使用与实验C相同的亮场、暗场数据，从光源减弱及弱信号可测性的角度进行分析，不作为另一组独立实验数据。光源通过降低供电电压逐级减弱，共设置60、63、68、72、76 V五档电压。")
table(["光源电压", "积分时间", "每个条件亮场/暗场帧数"], [
    ["60 V", "64、128 ms", "30/30"], ["63 V", "64、128 ms", "30/30"],
    ["68 V", "32、64 ms", "30/30"], ["72 V", "32、64 ms", "30/30"],
    ["76 V", "32、64 ms", "32 ms：50/50；64 ms：30/30"],
], widths=[3.5, 4.5, 7.5], font_size=9.2)
p("全部电压档均包含64 ms数据，可用于同积分时间条件下的横向比较。数据覆盖288个通道，波长范围为305.883–878.092 nm。")
p("本次未记录参考板辐亮度，也未使用已知透过率的中性密度片。因此：")
bullet("电压仅用于标识光源设置，不能直接换算为辐亮度；")
bullet("不能将电压比视为光强衰减比例；")
bullet("不能假设降低电压时各波长信号按相同比例减弱。")

doc.add_heading("D.3 评价方法", level=2)
doc.add_heading("D.3.1 扣暗与重复性统计", level=3)
p("对每个电压和积分时间条件，使用同条件平均暗谱进行扣暗：")
equation("D̄(λ) = (1/nD) Σ Dⱼ(λ)，Sᵢ(λ) = Lᵢ(λ) − D̄(λ)")
p("L、D导出值中均包含1000 DN固定偏置，相减后自动抵消，不再额外扣除。逐波长净信号均值、样本标准差、信噪比与变异系数为：")
equation("S̄(λ) = (1/nL) Σ Sᵢ(λ)")
equation("sS(λ) = √{Σ[Sᵢ(λ)−S̄(λ)]²/(nL−1)}")
equation("SNR(λ) = S̄(λ)/sS(λ)，CV(λ) = sS(λ)/|S̄(λ)| × 100%")
p("对于正信号，CV=100%/SNR。因此，SNR≥20与CV≤5%是等价要求；SNR≥50与CV≤2%是等价的推荐要求。")

doc.add_heading("D.3.2 最低可测亮度的定义", level=3)
p("以Lref表示参考板实测辐亮度，以Λout表示预先确定的正式输出波段。最低可测亮度Lmin定义为：在明确的积分时间及采集方式下，正式输出波段内全部通道均满足净光均值为正、SNR≥20、CV≤5%且无饱和时，对应的最低实测参考板辐亮度。推荐工作区应满足SNR≥50、CV≤2%。实际参数选择还需结合80%量程余量及采集期漂移≤2%的要求。")
p("Lmin与输出波段、积分时间和帧平均方式有关，不能脱离这些条件单独给出。本次没有辐亮度记录，不能计算具有物理单位的Lmin。")

doc.add_heading("D.3.3 延长积分时间的效果", level=3)
equation("G(λ) = SNR长(λ)/SNR短(λ)")
equation("Q(λ) = [S̄长(λ)/t长] / [S̄短(λ)/t短]")
p("G>1表示延长积分时间后SNR提高；Q接近1表示两档单位积分时间响应接近。上述比较仍可能受到两次采集之间光源变化的影响。")

doc.add_heading("D.4 实验结果", level=2)
doc.add_heading("D.4.1 固定积分时间下的响应变化", level=3)
p("固定积分时间为64 ms，统计500–800 nm中段通道的净光响应与SNR中位数：")
table(["光源电压", "中段净光响应中位数 / DN", "中段SNR中位数"], [
    ["60 V", "12864", "122.7"], ["63 V", "17337", "140.4"],
    ["68 V", "25785", "168.9"], ["72 V", "34717", "210.8"],
    ["76 V", "44319", "262.1"],
], widths=[3.5, 7.0, 5.0], font_size=9.2)
p("随着电压降低，中段净光响应与SNR均下降。最低已测电压档60 V、64 ms下，中段SNR中位数仍为122.7，高于基本要求20及推荐要求50。")
p("但是，中段中位数不能代表全部通道。60 V、64 ms条件下仍有53个通道SNR低于20，主要集中于短波端。因此，本次结果支持“最低已测电压档下中段信号仍具有较高SNR”，不能表述为“最低档下全波段均满足测量要求”。此外，76 V、64 ms存在部分通道超过80%量程，表中的高SNR不代表该条件可直接作为全波段工作设置。")

doc.add_heading("D.4.2 降压对光谱形状的影响", level=3)
p("在相同64 ms积分时间下，60 V与72 V数据的谱形比较结果为：谱形角约6.4°，归一化差异指标NRMSE约11.7%。该结果表明，不同电压档下的响应不能简单视为同一光谱乘以不同衰减系数。弱光分析应保留逐波长结果，不能用单一电压比例或全谱平均信号比例替代各波长的实际衰减。")
p("由于缺少独立光源光谱记录，上述差异不能全部归因于光源，也可能包含仪器响应及测量波动的影响。")

doc.add_heading("D.4.3 延长积分时间的改善效果", level=3)
table(["电压", "积分时间变化", "SNR增益G中位数", "SNR≥20通道数变化", "单位时间响应比Q中位数"], [
    ["60 V", "64→128 ms", "1.516", "235→264", "1.00515"],
    ["63 V", "64→128 ms", "1.527", "247→285", "1.00210"],
    ["68 V", "32→64 ms", "1.508", "238→275", "1.00617"],
    ["72 V", "32→64 ms", "1.510", "260→288", "1.00430"],
    ["76 V", "32→64 ms", "1.554", "271→288", "0.99777"],
], widths=[2.2, 3.3, 3.2, 3.6, 3.5], font_size=8.0)
p("积分时间加倍后，逐波长SNR增益中位数为1.51–1.55，说明在本次条件下，延长积分时间能够改善弱信号重复性。在最低已测电压60 V下，积分时间由64 ms延长至128 ms后，SNR不达标通道由53个减少至24个。两档单位时间响应比中位数接近1，表明净光响应总体接近按积分时间比例增加。")
p("但延长积分时间并非始终有利。76 V下，64 ms虽然提高了SNR，却使78个通道超过80%量程，工作候选通道数反而由32 ms下的264个减少至199个。因此，积分时间需同时兼顾弱信号SNR和强信号码值余量。")

doc.add_heading("D.4.4 最低已测电压档的可用波段", level=3)
table(["指标", "60 V / 64 ms", "60 V / 128 ms"], [
    ["SNR≥20通道数", "235", "264"], ["超过80%量程的通道数", "0", "0"],
    ["漂移超过2%的通道数", "49", "39"], ["工作候选通道数", "224", "241"],
    ["推荐候选通道数", "205", "223"], ["最长连续工作候选波段", "499.1–837.6 nm", "474.6–878.1 nm"],
    ["最长连续推荐候选波段", "515.9–837.6 nm", "474.6–878.1 nm"],
], widths=[6.5, 4.5, 4.5], font_size=8.8)
p("60 V、128 ms条件下，474.6–878.1 nm内的223个连续通道同时满足推荐SNR、码值余量与漂移要求，说明延长积分时间能够扩大最低已测电压档下的连续可用波段。")
p("该结果描述的是“60 V这一已测设置下的表现”，不能直接换算为最低可测辐亮度，也不能证明更低电压或更低亮度下仍可保持相同性能。")

doc.add_heading("D.5 弱光趋势的探索性分析", level=2)
p("在固定64 ms条件下，使用五档电压对应的500–800 nm净光响应中位数与SNR中位数进行经验拟合：")
equation("SNRmed = A(S̄med)ᵖ，p ≈ 0.601")
p("按该经验关系外推，SNR中位数降至20时，对应的净光响应约为669 DN，为最低已测点12864 DN的约1/19.2。")
p("这一结果仅为探索性数学外推，不作为最低可测亮度或实际弱光余量结论。", bold_lead="这一结果仅为探索性数学外推，不作为最低可测亮度或实际弱光余量结论。")
numbered("实测点仅有5个，响应范围为12864–44319 DN，外推位置远低于全部实测点。")
numbered("使用的是波段中位数，不能保证各通道均达到同一SNR要求。")
numbered("不同电压档的谱形存在差异，低亮度下是否仍服从同一经验关系未经验证。")
numbered("更低信号下的噪声组成、暗场稳定性及漂移影响可能变化。")
p("拟合指数也不能单独用于判定散粒噪声、暗电流噪声或读出噪声中的哪一种占主导。")
figure("06_lowlight_scaling.png", "图6  固定64 ms条件下中段净光响应与SNR的经验关系。散点为五档电压下500–800 nm通道的净光响应中位数与SNR中位数；实测范围外部分属于探索性外推，不用于确定最低可测亮度。", width_cm=13.8)

doc.add_heading("D.6 参数建议与改进需求", level=2)
doc.add_heading("D.6.1 现有数据支持的参数选择", level=3)
p("在本次最低已测电压60 V下，128 ms相较64 ms具有更高SNR和更宽的连续工作候选波段，可作为类似室内弱信号条件下的优先试验设置。")
p("本次信噪比与弱光数据仅覆盖32、64、128 ms三个积分时间水平。128 ms是本次已测试的最长积分时间，不等同于已经确定的推荐积分时间上限。更长积分时间仍需检查任务时效、强波段量程、暗场变化及光源漂移。")
p("本设备增益固定，不提供可调增益推荐值。本次统计基于单帧序列的离散程度，尚未验证不同平均帧数的实际收益，因此不指定统一的最优平均帧数。")

doc.add_heading("D.6.2 无效波段判定", level=3)
bullet("净光均值非正，或SNR＜20：标记为信号不足；")
bullet("亮场扣除固定偏置后超过80%量程：标记为码值余量不足；")
bullet("拟合漂移绝对值＞2%：标记为采集期稳定性未通过；")
bullet("存在数字饱和或数据异常：标记为无效。")
p("仅通过全部工作要求的通道纳入工作候选波段。无效标记针对具体采集条件，不表示该波长在其他参数下永久不可测。")

doc.add_heading("D.6.3 后续验证与改进方向", level=3)
table(["待确认问题", "验证方式", "根据验证结果采取的措施"], [
    ["弱信号不足", "在更低实测辐亮度下比较不同积分时间", "优先评估增加有效通光量或延长积分时间，并同步检查强波段量程"],
    ["暗场贡献增大", "比较不同积分时间、温度下的暗场均值与波动", "若确认暗电流相关影响显著，再评估温控及积分时间限制"],
    ["读出噪声影响", "结合暗场与不同信号水平的数据分析噪声组成", "本设备无可调增益，可评估多帧平均及采集链路优化"],
    ["长序列漂移", "按固定条件采集不少于100帧并记录时间戳", "根据漂移来源改善光源稳定性、温控或采集时序"],
    ["杂散光或背景干扰", "开展遮光、背景及光路对照测量", "若确认存在影响，再改善遮光结构及光路"],
    ["多帧平均收益", "比较不同平均帧数下的重复性与耗时", "在精度与任务时间之间确定平均帧数"],
], widths=[3.2, 5.3, 7.0], font_size=7.8)
p("以上为后续验证方向，现有数据不足以确定哪一种噪声机制主导，也不足以直接认定必须进行某项硬件改造。")

doc.add_heading("D.7 结论与适用范围", level=2)
numbered("最低已测电压档下，中段信号仍具有较高SNR，但短波端存在不足。60 V、64 ms条件下，500–800 nm中段SNR中位数为122.7，但全谱仍有53个通道SNR低于20。")
numbered("延长积分时间能够改善已测弱信号条件下的性能。60 V下由64 ms延长至128 ms后，SNR不达标通道减少至24个，最长连续工作候选波段由499.1–837.6 nm扩展至474.6–878.1 nm。")
numbered("本次不能给出绝对最低可测亮度。缺少参考板辐亮度记录，且仅采集五档电压，未证明覆盖工程最低亮度以下的条件，因此Lmin尚未确定。")
numbered("推荐积分时间上限与最优平均帧数尚需进一步验证。本次最长测试积分时间为128 ms，采集序列为30或50帧，不能替代更长积分时间及100帧稳定性试验。")
numbered("现有结果可用于室外实验的初始参数选择。室外目标的光谱分布、背景、温度及时间变化与室内不同，最终可用波段和参数应通过现场亮场、暗场及重复测量数据确定。")


# 页码
for section in doc.sections:
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    para.add_run(" 页")
    for rr in para.runs:
        rr.font.size = Pt(9)
        rr.font.name = "宋体"
        rr._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

# 文档属性
doc.core_properties.title = "二代 OCR 模组室内实验分析报告"
doc.core_properties.subject = "实验 B 积分时间线性度 实验 C 信噪比 实验 D 弱光性能"
doc.core_properties.author = ""
doc.core_properties.keywords = "OCR, 光谱仪, 线性度, 信噪比, 弱光性能"

doc.save(OUT)
print(OUT)
