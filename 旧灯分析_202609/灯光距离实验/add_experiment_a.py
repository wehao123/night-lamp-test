from pathlib import Path
import base64
import io

from lxml import html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\haowei\Documents\ChatGPT\灯光距离实验")
SOURCE = ROOT / "OCR-25F02100_室内实验B_C_D完整报告.docx"
HTML = ROOT / "outputs" / "lamp_warmup_20260922" / "分析报告.html"
OUTPUT = ROOT / "OCR-25F02100_室内实验A_B_C_D完整报告.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, value=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_rich_text(paragraph, element):
    if element.text:
        paragraph.add_run(element.text)
    for child in element:
        text = "".join(child.itertext())
        run = paragraph.add_run(text)
        if child.tag.lower().endswith(("b", "strong")):
            run.bold = True
        if child.tail:
            paragraph.add_run(child.tail)


def add_table(doc, table_el):
    rows = table_el.xpath(".//tr")
    if not rows:
        return
    headers = [" ".join(x.itertext()).strip() for x in rows[0].xpath("./th|./td")]
    body = [[" ".join(x.itertext()).strip() for x in row.xpath("./th|./td")] for row in rows[1:]]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    repeat_header(table.rows[0])
    for j, value in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = value
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(8.5)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        cell_margins(cell)
    for i, values in enumerate(body):
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = value
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cells[j].paragraphs[0].paragraph_format.space_after = Pt(0)
            for run in cells[j].paragraphs[0].runs:
                run.font.size = Pt(8.5)
                run.font.name = "宋体"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            if i % 2:
                shade(cells[j], "EEF4F8")
            cell_margins(cells[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


doc = Document(SOURCE)
before_count = len(doc.element.body)

# 实验 A 总标题
doc.add_heading("实验 A  灯光冷启动稳定性", level=1)

tree = html.fromstring(HTML.read_text(encoding="utf-8"))
main = tree.xpath("//main")[0]
h2_index = 1
image_index = 1

for node in main:
    tag = node.tag.lower() if isinstance(node.tag, str) else ""
    if tag == "h1":
        doc.add_heading(f"A.{h2_index} {node.text_content().strip()}", level=2)
        h2_index += 1
    elif tag == "h2":
        doc.add_heading(f"A.{h2_index} {node.text_content().strip()}", level=2)
        h2_index += 1
    elif tag == "p":
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0.74)
        add_rich_text(para, node)
    elif tag == "div" and "table" in (node.get("class") or ""):
        tables = node.xpath(".//table")
        if tables:
            add_table(doc, tables[0])
    elif tag == "table":
        add_table(doc, node)
    elif tag == "img":
        src = node.get("src", "")
        if src.startswith("data:image/") and ";base64," in src:
            raw = base64.b64decode(src.split(",", 1)[1])
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(io.BytesIO(raw), width=Cm(15.8))
            cap = doc.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.add_run(f"图A-{image_index}  {node.get('alt', '实验结果图')}")
            image_index += 1
    elif tag in ("ul", "ol"):
        style = "List Bullet" if tag == "ul" else "List Number"
        for li in node.xpath("./li"):
            para = doc.add_paragraph(style=style)
            para.paragraph_format.left_indent = Cm(0.75)
            para.add_run(li.text_content().strip())

# A 与 B 之间分页
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# 把刚创建的 A 章节元素移动到实验 B 之前
body = doc.element.body
new_elements = list(body)[before_count:]
target = None
for para in doc.paragraphs:
    if para.text.strip().startswith("实验 B"):
        target = para._p
        break
if target is None:
    raise RuntimeError("未找到实验 B 标题")
for element in new_elements:
    target.addprevious(element)

# 更新封面说明
for para in doc.paragraphs:
    if "实验 B 积分时间线性度" in para.text and "实验 D 弱光性能" in para.text:
        para.text = "实验 A 灯光冷启动稳定性  实验 B 积分时间线性度  实验 C 信噪比  实验 D 弱光性能"
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.name = "黑体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            run.font.size = Pt(12)
    if para.text.startswith("本报告汇总积分时间相对线性度"):
        para.text = "本报告汇总灯光冷启动稳定性、积分时间相对线性度、信噪比和弱光性能四项室内实验结果。全部定量结果均由原始重复帧计算，未进行平滑、滤波或插值。"

doc.core_properties.subject = "实验 A 灯光冷启动稳定性 实验 B 积分时间线性度 实验 C 信噪比 实验 D 弱光性能"
doc.save(OUTPUT)
print(OUTPUT)
