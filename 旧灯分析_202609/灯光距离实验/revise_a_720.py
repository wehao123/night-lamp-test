from pathlib import Path
import json, csv, html, io
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from docx import Document

BASE=Path(__file__).resolve().parent
ROOT=BASE/'outputs/lamp_warmup_20260922'
OUT=ROOT/'环境1完整30帧_红光至720nm'
OUT.mkdir(exist_ok=True)
# Reuse the existing source reader and independent numerical calculations only.
source=(ROOT/'audit_three_backgrounds.py').read_text(encoding='utf-8').split('checks=[]')[0]
source=source.replace("OUT=ROOT/'三组背景重算'","OUT=ROOT/'环境1完整30帧_红光至720nm'")
source=source.replace('(420,700)','(420,720)').replace('(600,700)','(600,720)')
ns={'__file__':str(ROOT/'audit_three_backgrounds.py')}
exec(compile(source,'recompute_720','exec'),ns)
res=ns['results']['固定环境1完整30帧']; y=ns['nets']['固定环境1完整30帧']; sec=ns['sec']; env=ns['env']; integ=ns['integ']
assert np.allclose(y[:,0],y[:,1:].sum(axis=1))
assert np.allclose(integ(ns['x'][[0,2000,-1]]-env[0]['x'].mean(axis=0)),y[[0,2000,-1]])
old=json.loads((ROOT/'环境1完整30帧重算/结果.json').read_text(encoding='utf-8'))
for j in [1,2]:
    assert abs(res['pp'][j]-old['pp'][j])<1e-9
(OUT/'结果.json').write_text(json.dumps(res,ensure_ascii=False,indent=2),encoding='utf-8')
with (OUT/'逐帧净积分.csv').open('w',encoding='utf-8-sig',newline='') as f:
    wr=csv.writer(f);wr.writerow(['相对首帧秒','总量420-720','蓝420-500','绿500-600','红600-720']);wr.writerows(zip(sec,*y.T))

# Fresh charts from the 720 nm data; do not reuse earlier 700 nm charts.
FONT='C:/Windows/Fonts/msyh.ttc'
def font(n):return ImageFont.truetype(FONT,n)
colors=['#263e56','#287dc1','#278047','#c74b49']
labels=['总量 420–720 nm','蓝 420–500 nm','绿 500–600 nm','红 600–720 nm']
def chart(filename,title,series,xlabel,ylabel,threshold=None):
    im=Image.new('RGB',(1700,900),'white');d=ImageDraw.Draw(im)
    d.text((100,20),title,font=font(34),fill='black');d.text((110,85),ylabel,font=font(25),fill='black')
    xs=[p[0] for _,pts,_ in series for p in pts if np.isfinite(p[1])];ys=[p[1] for _,pts,_ in series for p in pts if np.isfinite(p[1])]
    if threshold is not None:ys.append(threshold)
    xmin,xmax=min(xs),max(xs); ymin,ymax=min(ys),max(ys);pad=(ymax-ymin)*.08 or 1;ymin-=pad;ymax+=pad
    X=lambda x:145+(x-xmin)/(xmax-xmin or 1)*1470
    Y=lambda v:715-(v-ymin)/(ymax-ymin)*575
    for k in range(6):
        v=ymin+(ymax-ymin)*k/5;xx=xmin+(xmax-xmin)*k/5
        d.line((145,Y(v),1615,Y(v)),fill='#dddddd',width=1)
        d.text((10,Y(v)-14),f'{v:.3f}' if abs(ymax)<2 else f'{v:.1f}',font=font(23),fill='black')
        d.text((X(xx)-25,735),f'{xx:.1f}',font=font(23),fill='black')
    if threshold is not None:
        d.line((145,Y(threshold),1615,Y(threshold)),fill='#999999',width=3)
        d.text((1350,Y(threshold)-32),f'判据 {threshold:g}',font=font(23),fill='#555555')
    for k,(label,pts,gap) in enumerate(series):
        prev=None
        for x,v in pts:
            if not np.isfinite(v):prev=None;continue
            cur=(X(x),Y(v))
            if prev is not None and x-prev[0]<=gap:d.line((prev[1],prev[2],*cur),fill=colors[k%4],width=3)
            d.ellipse((cur[0]-2,cur[1]-2,cur[0]+2,cur[1]+2),fill=colors[k%4])
            prev=(x,*cur)
        lx=145+k*385;d.line((lx,845,lx+40,845),fill=colors[k%4],width=4);d.text((lx+48,829),label,font=font(22),fill=colors[k%4])
    d.text((650,780),xlabel,font=font(25),fill='black')
    im.save(OUT/filename,dpi=(240,240))

series=[]
for j in range(4):
    pts=[];ref=y[sec<10,j].mean()
    for b in range(int(sec[-1]//10)+1):
        m=(sec>=b*10)&(sec<(b+1)*10);tt=sec[m]
        valid=len(tt)>=10 and tt[0]-b*10<=1 and (b+1)*10-tt[-1]<=1 and np.max(np.diff(tt))<=2
        pts.append(((b+.5)/6,100*(y[m,j].mean()/ref-1) if valid else np.nan))
    series.append((labels[j],pts,.2))
chart('01_全程响应变化.png','固定环境1完整30帧背景后的响应变化',series,'相对首帧时间 / min','相对首10秒变化 / %')
rolling=[]
for e in range(300,int(sec[-1])+1):
    m=(sec>=e-300)&(sec<=e);tt=sec[m];yy=y[m,0]
    if len(tt)<10 or tt[0]-(e-300)>1 or e-tt[-1]>1 or np.max(np.diff(tt))>2:continue
    cv=yy.std(ddof=1)/yy.mean()*100
    rate=np.dot(tt-tt.mean(),yy-yy.mean())/np.sum((tt-tt.mean())**2)/yy.mean()*6000
    rolling.append((e/60,cv,abs(rate)))
chart('cv.png','总量420–720 nm 五分钟滑动CV',[('总量',[(t,c) for t,c,r in rolling],.018)],'窗口结束时刻 / min','CV / %',1)
chart('rate.png','总量420–720 nm 五分钟滑动变化率',[('总量',[(t,r) for t,c,r in rolling],.018)],'窗口结束时刻 / min','变化率绝对值 / (%/min)',.2)
im=Image.new('RGB',(1700,1800),'white');im.paste(Image.open(OUT/'cv.png'),(0,0));im.paste(Image.open(OUT/'rate.png'),(0,900));im.save(OUT/'02_滑动稳定性.png',dpi=(240,240))
envq=[integ(f['x']-1000) for f in env]
chart('03_环境背景对照.png','三组环境背景均值对照（已去除1000 DN偏置）',[(labels[j],[(k+1,float(q[:,j].mean()/1e5)) for k,q in enumerate(envq)],2) for j in range(4)],'环境背景组号','积分响应 / (10⁵ DN·nm)')
chart('04_环境背景逐帧.png','三组环境背景逐帧变化（420–720 nm）',[(f'环境{k+1}，{len(q)}帧',list(zip(range(1,len(q)+1),q[:,0]/1e5)),2) for k,q in enumerate(envq)],'各组内部帧序号','去偏置积分响应 / (10⁵ DN·nm)')

doc=Document(BASE/'OCR-25F02100_室内实验A完整报告.docx')
mapping={'420—700':'420—720','600—700':'600—720'}
for key in ['pp','cv','endpoint']:
    for a,b in zip(old[key],res[key]):
        mapping[f'{a:.3f}']=f'{b:.3f}';mapping[f'{abs(a):.3f}']=f'{abs(b):.3f}'
assert res['stable'] is not None
for key in ['max_cv','max_rate']:mapping[f"{old['stable'][key]:.3f}"]=f"{res['stable'][key]:.3f}"
for key in ['start_min','confirmed_min']:mapping[f"{old['stable'][key]:.2f}"]=f"{res['stable'][key]:.2f}"
import re
pattern=re.compile('|'.join(re.escape(k) for k in sorted(mapping,key=len,reverse=True)))
pars=list(doc.paragraphs)+[p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
for p in pars:
    for run in p.runs:
        if run.text:run.text=pattern.sub(lambda m:mapping[m.group()],run.text)
for name,shape in zip(['01_全程响应变化.png','02_滑动稳定性.png','03_环境背景对照.png','04_环境背景逐帧.png'],doc.inline_shapes):
    rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    doc.part.related_parts[rid]._blob=(OUT/name).read_bytes()
    image=Image.open(OUT/name);shape.height=int(shape.width*image.height/image.width)
for p in doc.paragraphs:
    if '不能直接作为时变背景插值使用' in p.text:p.text=p.text.replace('不能直接作为时变背景插值使用','本版仅作背景对照，不用于时变背景插值')
    if p.text.startswith('图A-3'):p.text='图A-3  三组环境背景总量及分波段平均积分响应对照，均已去除1000 DN偏置。'
    if p.text.startswith('图A-4'):p.text='图A-4  三组环境背景420—720 nm逐帧积分响应，横轴为各组内部帧序号。'
target=BASE/'OCR-25F02100_室内实验A完整报告_红光至720nm.docx'
doc.save(target)
check=Document(target);text='\n'.join(p.text for p in check.paragraphs)+'\n'+'\n'.join(c.text for t in check.tables for row in t.rows for c in row.cells)
assert '420—700' not in text and '600—700' not in text and len(check.inline_shapes)==4
assert f"{res['pp'][0]:.3f}" in text and f"{res['pp'][3]:.3f}" in text
(OUT/'修订记录.txt').write_text('红光600–720 nm，总量420–720 nm；环境1完整30帧固定背景。重新积分6899帧，蓝绿结果一致，总量等于三分波段之和；4张图均为本次新生成。\n'+json.dumps(res,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(res,ensure_ascii=False,indent=2));print(target)
