from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
targets=[
    ROOT/'室内分项实验综合报告_420–720nm修订版_v3.docx',
    ROOT/'OCR-25F02100_室内实验B_C_D完整报告_公式修订.docx',
    ROOT/'OCR-25F02100_室内实验A_B_C_D完整报告_公式修订.docx',
]
repl={
    'Sᵢ(λ) = a(λ) + k(λ)τᵢ + εᵢ':'Sᵢ(λ) = a(λ) + k(λ)τᵢ + εᵢ',
    'Dfit(λ) = k(λ)(τ末−τ首)/|S̄(λ)| × 100%':'D_fit(λ) = k(λ)[τ_end − τ_start] / |S̄(λ)| × 100%',
}
for path in targets:
    if not path.exists(): continue
    doc=Document(path); changed=0
    for p in doc.paragraphs:
        old=p.text
        if old in repl:
            for r in list(p.runs): r._element.getparent().remove(r._element)
            r=p.add_run(repl[old]); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r.font.name='Cambria Math'; r.font.size=Pt(11); changed+=1
        elif '根据实际时间戳，对各通道净信号进行线性拟合' in p.text:
            p.text=p.text.replace('根据实际时间戳，对各通道净信号进行线性拟合：','根据实际时间戳，对各通道净信号进行线性拟合。τ_start和τ_end分别表示统计区间的起始与结束时刻：')
            changed+=1
    if changed:
        out=path.with_name(path.stem+'_漂移公式修订.docx'); doc.save(out); print(out.name,changed)
