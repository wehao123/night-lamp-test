from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import numpy as np
import openpyxl

ROOT = Path(r"C:\Users\haowei\Documents\ChatGPT\灯光距离实验")
BOOK = ROOT / "outputs" / "01a09f15-4651-76a2-9e41-23c8b8d0b1e5" / "LED_distance_spot_analysis.xlsx"
OUT = ROOT / "figures"
OUT.mkdir(exist_ok=True)
wb = openpyxl.load_workbook(BOOK, data_only=True, read_only=True)
FONT_PATH = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD_PATH = r"C:\Windows\Fonts\msyhbd.ttc"
COLORS = ["#2F5597", "#ED7D31", "#70AD47", "#A64D79", "#5B9BD5", "#C00000"]


def font(size, bold=False):
    path = FONT_BOLD_PATH if bold and Path(FONT_BOLD_PATH).exists() else FONT_PATH
    return ImageFont.truetype(path, size)


class Chart:
    def __init__(self, width=1400, height=820, title="", xlabel="", ylabel="", margin=None):
        self.im = Image.new("RGB", (width, height), "white")
        self.d = ImageDraw.Draw(self.im)
        self.w, self.h = width, height
        self.m = margin or (145, 92, 55, 120)
        self.title, self.xlabel, self.ylabel = title, xlabel, ylabel

    def bounds(self):
        l, t, r, b = self.m
        return l, t, self.w - r, self.h - b

    def axes(self, xmin, xmax, ymin, ymax, xticks, yticks, xfmt=str, yfmt=str):
        l, t, r, b = self.bounds()
        self.d.line((l, b, r, b), fill="#333333", width=2)
        self.d.line((l, t, l, b), fill="#333333", width=2)
        f = font(24)
        for y in yticks:
            py = b - (y - ymin) / (ymax - ymin) * (b - t)
            self.d.line((l, py, r, py), fill="#D9D9D9", width=1)
            txt = yfmt(y)
            box = self.d.textbbox((0, 0), txt, font=f)
            self.d.text((l - 14 - (box[2] - box[0]), py - 14), txt, fill="#444444", font=f)
        for x in xticks:
            px = l + (x - xmin) / (xmax - xmin) * (r - l)
            self.d.line((px, b, px, b + 7), fill="#333333", width=2)
            txt = xfmt(x)
            box = self.d.textbbox((0, 0), txt, font=f)
            self.d.text((px - (box[2] - box[0]) / 2, b + 12), txt, fill="#444444", font=f)
        self.d.text((self.w / 2, 25), self.title, anchor="ma", fill="#222222", font=font(34, True))
        self.d.text((self.w / 2, self.h - 35), self.xlabel, anchor="ms", fill="#333333", font=font(26))
        yf = font(25)
        box = self.d.textbbox((0, 0), self.ylabel, font=yf)
        label = Image.new("RGBA", (box[2] - box[0] + 14, box[3] - box[1] + 14), (255, 255, 255, 0))
        ImageDraw.Draw(label).text((7, 7), self.ylabel, fill="#333333", font=yf, anchor="la")
        label = label.rotate(90, expand=True)
        self.im.paste(label, (18, int((t + b - label.height) / 2)), label)
        self.scale = (xmin, xmax, ymin, ymax)

    def xy(self, x, y):
        l, t, r, b = self.bounds()
        xmin, xmax, ymin, ymax = self.scale
        return l + (x - xmin) / (xmax - xmin) * (r - l), b - (y - ymin) / (ymax - ymin) * (b - t)

    def line(self, xs, ys, color, width=4, dots=True, dashed=False):
        pts = [self.xy(float(x), float(y)) for x, y in zip(xs, ys)]
        if dashed:
            for a, b in zip(pts[:-1], pts[1:]):
                dx, dy = b[0] - a[0], b[1] - a[1]
                length = max(1, (dx * dx + dy * dy) ** 0.5)
                for s in np.arange(0, length, 18):
                    e = min(s + 10, length)
                    p1 = (a[0] + dx * s / length, a[1] + dy * s / length)
                    p2 = (a[0] + dx * e / length, a[1] + dy * e / length)
                    self.d.line((p1, p2), fill=color, width=width)
        else:
            self.d.line(pts, fill=color, width=width)
        if dots:
            for px, py in pts:
                self.d.ellipse((px - 6, py - 6, px + 6, py + 6), fill=color, outline="white", width=2)

    def hline(self, y, color="#777777"):
        l, _, r, _ = self.bounds()
        py = self.xy(self.scale[0], y)[1]
        for x in range(int(l), int(r), 24):
            self.d.line((x, py, min(x + 14, r), py), fill=color, width=2)

    def legend(self, items, x=None, y=None, cols=1):
        x = x or self.bounds()[0] + 15
        y = y or self.bounds()[1] + 15
        f = font(21)
        colw = 245
        for i, (label, color, dashed) in enumerate(items):
            row, col = divmod(i, cols)
            xx, yy = x + col * colw, y + row * 35
            if dashed:
                for sx in range(int(xx), int(xx + 42), 14):
                    self.d.line((sx, yy + 11, sx + 8, yy + 11), fill=color, width=3)
            else:
                self.d.line((xx, yy + 11, xx + 42, yy + 11), fill=color, width=4)
            self.d.text((xx + 52, yy), label, fill="#333333", font=f)

    def save(self, name):
        self.im.save(OUT / name)


ws = wb["Summary"]
distance = np.array([ws.cell(r, 1).value for r in range(18, 23)], float)
response = np.array([ws.cell(r, 2).value for r in range(18, 23)], float)
fit_response = np.array([ws.cell(r, 4).value for r in range(18, 23)], float)
relative = response / response[0] * 100
c = Chart(title="中心净响应随距离衰减", xlabel="距离 (m)", ylabel="相对净响应 (%)")
c.axes(1.4, 5.15, 0, 112, distance, [0, 20, 40, 60, 80, 100], lambda x: f"{x:g}", lambda y: f"{int(y)}")
c.line(distance, relative, COLORS[0])
c.line(distance, fit_response / response[0] * 100, COLORS[1], dashed=True, dots=False)
for x, y in zip(distance, relative):
    px, py = c.xy(x, y)
    c.d.text((px, py - 18), f"{y:.1f}%", anchor="ms", font=font(20), fill="#333333")
c.legend([("实测中心响应", COLORS[0], False), ("幂律拟合 R∝r^-1.88", COLORS[1], True)], x=760, y=120)
c.save("01_distance_response.png")

xpos = np.array([ws.cell(r, 1).value for r in range(27, 32)], float)
labels = ["1.507 m", "2.032 m", "3.016 m", "3.995 m", "5.013 m"]
profiles = np.array([[ws.cell(r, col).value for r in range(27, 32)] for col in range(2, 7)], float) * 100
c = Chart(title="不同距离下的横向归一化光斑", xlabel="横向位置 (cm)", ylabel="相对中心响应 (%)")
c.axes(-42, 42, 0, 108, [-40, -20, 0, 20, 40], [0, 20, 40, 60, 80, 100], lambda x: f"{int(x)}", lambda y: f"{int(y)}")
for vals, color in zip(profiles, COLORS):
    c.line(xpos, vals, color)
c.hline(80)
c.legend([(l, co, False) for l, co in zip(labels, COLORS)] + [("80%阈值", "#777777", True)], x=175, y=120, cols=2)
c.save("02_lateral_profiles.png")

udist = np.array([ws.cell(r, 1).value for r in range(37, 42)], float)
min20 = np.array([ws.cell(r, 2).value for r in range(37, 42)], float) * 100
min25 = np.array([ws.cell(r, 4).value for r in range(37, 42)], float) * 100
width80 = np.array([ws.cell(r, 5).value for r in range(37, 42)], float)
im = Image.new("RGB", (1800, 760), "white")
left = Chart(900, 760, "中心区域均匀性", "距离 (m)", "区域内最低相对响应 (%)", margin=(135, 90, 35, 120))
left.axes(1.4, 5.15, 0, 105, udist, [0, 20, 40, 60, 80, 100], lambda x: f"{x:g}", lambda y: f"{int(y)}")
left.line(udist, min20, COLORS[0]); left.line(udist, min25, COLORS[1]); left.hline(80)
left.legend([("±20 cm最低值", COLORS[0], False), ("±25 cm估算最低值", COLORS[1], False)], x=160, y=115)
right = Chart(900, 760, "80%均匀光斑宽度", "距离 (m)", "估算宽度 (cm)", margin=(125, 90, 45, 120))
right.axes(0, 5, 0, 66, [.5, 1.5, 2.5, 3.5, 4.5], [0, 10, 20, 30, 40, 50, 60], lambda x: labels[int(x-.5)].split()[0], lambda y: f"{int(y)}")
for i, (y, color) in enumerate(zip(width80, COLORS)):
    x0, y0 = right.xy(i+.18, 0); x1, y1 = right.xy(i+.82, y)
    right.d.rectangle((x0, y1, x1, y0), fill=color)
    right.d.text(((x0+x1)/2, y1-10), f"{y:.1f}", anchor="ms", font=font(20), fill="#333333")
right.hline(50, COLORS[1])
im.paste(left.im, (0, 0)); im.paste(right.im, (900, 0)); im.save(OUT / "03_uniformity_width.png")

asym20 = np.array([ws.cell(r, 7).value for r in range(37, 42)], float) * 100
asym40 = np.array([ws.cell(r, 8).value for r in range(37, 42)], float) * 100
repeat = np.array([ws.cell(r, 9).value for r in range(37, 42)], float) * 100
im = Image.new("RGB", (1800, 760), "white")
left = Chart(900, 760, "中心位置重复性", "距离 (m)", "两次中心测量差异 (%)", margin=(125, 90, 40, 120))
left.axes(0, 5, 0, 4.5, [.5, 1.5, 2.5, 3.5, 4.5], [0, 1, 2, 3, 4], lambda x: labels[int(x-.5)].split()[0], lambda y: f"{int(y)}")
for i, y in enumerate(repeat):
    x0, y0 = left.xy(i+.18, 0); x1, y1 = left.xy(i+.82, y)
    left.d.rectangle((x0, y1, x1, y0), fill="#70AD47")
    left.d.text(((x0+x1)/2, y1-9), f"{y:.1f}%", anchor="ms", font=font(20), fill="#333333")
right = Chart(900, 760, "左右不对称性", "距离 (m)", "不对称性 (%)", margin=(125, 90, 40, 120))
right.axes(0, 5, 0, 85, [.5, 1.5, 2.5, 3.5, 4.5], [0, 20, 40, 60, 80], lambda x: labels[int(x-.5)].split()[0], lambda y: f"{int(y)}")
for i, (a, b) in enumerate(zip(asym20, asym40)):
    for lo, hi, y, color in [(.10, .45, a, COLORS[0]), (.55, .90, b, COLORS[5])]:
        x0, y0 = right.xy(i+lo, 0); x1, y1 = right.xy(i+hi, y)
        right.d.rectangle((x0, y1, x1, y0), fill=color)
right.legend([("±20 cm", COLORS[0], False), ("±40 cm", COLORS[5], False)], x=585, y=115)
im.paste(left.im, (0, 0)); im.paste(right.im, (900, 0))
top = ImageDraw.Draw(im)
top.rectangle((0, 0, 1800, 78), fill="white")
top.text((450, 25), "中心位置重复性", anchor="ma", fill="#222222", font=font(34, True))
top.text((1350, 25), "左右不对称性", anchor="ma", fill="#222222", font=font(34, True))
im.save(OUT / "04_repeatability_symmetry.png")

bg = wb["Background"]
order = np.array([bg.cell(r, 1).value for r in range(7, 38)], int)
bg_dn = np.array([bg.cell(r, 5).value for r in range(7, 38)], float) / 1000
c = Chart(width=1550, title="遮阴背景随测量顺序变化", xlabel="测量顺序", ylabel="遮阴背景总DN (×10³)")
c.axes(1, 31, 275, 360, [1, 6, 7, 8, 13, 19, 25, 31], [280, 300, 320, 340, 360], lambda x: f"{int(x)}", lambda y: f"{int(y)}")
c.line(order, bg_dn, "#8064A2", width=3)
for boundary in [6.5, 7.5, 13.5, 19.5, 25.5]:
    px, _ = c.xy(boundary, 275)
    c.d.line((px, c.bounds()[1], px, c.bounds()[3]), fill="#BBBBBB", width=2)
c.d.text((c.xy(7.4, 352)[0], c.xy(7.4, 352)[1]), "约40.18 min间隔", anchor="mm", font=font(23), fill="#333333")
c.save("05_background_change.png")

spec = wb["Spectral analysis"]
wl = np.array([spec.cell(r, 1).value for r in range(7, 295)], float)
valid = np.array([spec.cell(r, 2).value for r in range(7, 295)], int) == 1
center_specs = np.array([[spec.cell(r, col).value for r in range(7, 295)] for col in range(3, 8)], float)
net = wb["Net spectra"]
near_headers = [net.cell(6, col).value for col in range(3, 9)]
near_specs = np.array([[net.cell(r, col).value for r in range(7, 295)] for col in range(3, 9)], float)

def normalized(arr):
    return arr / np.max(arr[:, valid], axis=1)[:, None]

im = Image.new("RGB", (1800, 760), "white")
left = Chart(900, 760, "中心净光谱形状", "波长 (nm)", "峰值归一化响应", margin=(115, 90, 35, 120))
left.axes(310, 855, 0, 1.08, [350, 450, 550, 650, 750, 850], [0, .2, .4, .6, .8, 1.0], lambda x: f"{int(x)}", lambda y: f"{y:.1f}")
for vals, color in zip(normalized(center_specs), COLORS):
    left.line(wl[valid], vals[valid], color, width=3, dots=False)
left.legend([(l, co, False) for l, co in zip(labels, COLORS)], x=135, y=112, cols=2)
right = Chart(900, 760, "1.507 m横向净谱：右2异常", "波长 (nm)", "峰值归一化响应", margin=(115, 90, 35, 120))
right.axes(310, 855, -.12, 1.08, [350, 450, 550, 650, 750, 850], [0, .2, .4, .6, .8, 1.0], lambda x: f"{int(x)}", lambda y: f"{y:.1f}")
for header, vals, color in zip(near_headers, normalized(near_specs), COLORS):
    right.line(wl[valid], vals[valid], "#C00000" if "右2" in header else color, width=5 if "右2" in header else 2, dots=False)
right.legend([(h.split(" ", 2)[-1], "#C00000" if "右2" in h else co, False) for h, co in zip(near_headers, COLORS)], x=135, y=112, cols=2)
im.paste(left.im, (0, 0)); im.paste(right.im, (900, 0))
top = ImageDraw.Draw(im)
top.rectangle((0, 0, 1800, 78), fill="white")
tf = font(34, True)
title_left = "中心净光谱形状"
title_right = "1.507 m横向净谱：右2异常"
left_box = top.textbbox((0, 0), title_left, font=tf)
right_box = top.textbbox((0, 0), title_right, font=tf)
top.text((450 - (left_box[2] - left_box[0]) / 2, 25), title_left, fill="#222222", font=tf)
top.text((1350 - (right_box[2] - right_box[0]) / 2, 25), title_right, fill="#222222", font=tf)
im.save(OUT / "06_spectral_shape.png")

print(f"Generated 6 figures in {OUT}")
