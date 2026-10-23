from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
targets=[
    ROOT/'室内分项实验综合报告_420–720nm修订版_v2.docx',
    ROOT/'室内分项实验综合报告_420–720nm修订版.docx',
    ROOT/'室内分项实验综合报告_按立项报告整理.docx',
    ROOT/'OCR-25F02100_室内实验B_C_D完整报告.docx',
    ROOT/'OCR-25F02100_室内实验A_B_C_D完整报告.docx',
]
repl={
    'S(λ,t) = K(λ)t':'S(λ,t) = K(λ) t',
    'S̄短(λ) = (1/n短) Σ Sᵢ(λ,t短)':'S̄ₛ(λ) = (1/nₛ) ∑ᵢ₌₁ⁿˢ Sᵢ(λ,tₛ)',
    'S̄长(λ) = (1/n长) Σ Sᵢ(λ,t长)':'S̄ₗ(λ) = (1/nₗ) ∑ᵢ₌₁ⁿˡ Sᵢ(λ,tₗ)',
    'S理论,长(λ) = S̄短(λ) · t长/t短':'Sₗⁱᵈᵉᵃˡ(λ) = S̄ₛ(λ) · (tₗ/tₛ)',
    'γ(λ) = S理论,长(λ) / S̄长(λ)':'γ(λ) = Sₗⁱᵈᵉᵃˡ(λ) / S̄ₗ(λ)',
    'δ(λ) = [S̄长(λ) − S理论,长(λ)] / S理论,长(λ) × 100%':'δ(λ) = [S̄ₗ(λ) − Sₗⁱᵈᵉᵃˡ(λ)] / Sₗⁱᵈᵉᵃˡ(λ) × 100%',
    'δ = (1/γ − 1) × 100%':'δ(λ) = [γ(λ)⁻¹ − 1] × 100%',
    'S̄短 > 0，S̄长 > 0，SNR短 ≥ 20，SNR长 ≥ 20':'S̄ₛ(λ) > 0,  S̄ₗ(λ) > 0,  SNRₛ(λ) ≥ 20,  SNRₗ(λ) ≥ 20',
}
def set_para(p,text,formula=False):
    for r in list(p.runs): r._element.getparent().remove(r._element)
    r=p.add_run(text)
    if formula:
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r.font.name='Cambria Math'; r.font.size=Pt(11)
        r._element.rPr.rFonts.set(__import__('docx').oxml.ns.qn('w:eastAsia'),'Cambria Math')
    return p
for path in targets:
    if not path.exists(): continue
    doc=Document(path); changed=0
    for p in doc.paragraphs:
        if p.text in repl:
            set_para(p,repl[p.text],True); changed+=1
        elif p.text.startswith('各档SNR由净光均值除以帧间样本标准差计算。'):
            p.text=p.text+' 本节的 S(λ,t)、γ(λ) 和 δ(λ) 为基于积分时间比例关系推导的计算定义；项目立项文件并未规定 γ 或 δ 的唯一命名方式。SNR≥20 是本报告采用的通道筛选门槛，95th percentile 和最大值属于本次样本的描述性统计，不是标准验收公式。'
            changed+=1
    if changed:
        out=path.with_name(path.stem+'_公式修订.docx')
        doc.save(out); print(out.name,changed)
