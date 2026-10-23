from pathlib import Path
from copy import deepcopy
import re, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

BASE=Path(__file__).resolve().parent
doc=Document(BASE/'source.docx')
body=doc.element.body
original=list(body)
paras=doc.paragraphs
table_names=[
'距离与光斑实验结果概览','不同距离的中心净响应','不同距离的横向归一化响应','横向最低响应与插值阈值宽度','横向不对称度与中心重复差','约40 min系统响应与背景变化',
'冷启动实验预处理规则','波段稳定性与全程首末变化','三次遮挡测量统计','背景处理方式敏感性比较',
'积分时间配对条件','积分响应统计指标','六组积分响应配对结果','净光码值分段统计',
'信噪比实验采集条件','逐通道联合筛选要求','各条件信噪比与联合筛选结果','延长积分时间的信噪比增益','最长连续候选波段',
'弱光分析共用数据条件','固定64 ms下的中段响应与信噪比','弱光条件下的积分时间对比','最低已测电压档的候选波段','弱光性能后续验证与改进方向']
tablemap={t._element:name for t,name in zip(doc.tables,table_names)}
starts={k:next(i for i,e in enumerate(original) if e.tag==qn('w:p') and ''.join(e.xpath('.//w:t/text()')).startswith(v)) for k,v in [('A','实验 A'),('B','实验 B'),('C','实验 C'),('D','实验 D')]}
sections={'distance':original[2:starts['A']], 'A':original[starts['A']:starts['B']], 'B':original[starts['B']:starts['C']], 'C':original[starts['C']:starts['D']], 'D':[e for e in original[starts['D']:] if e.tag!=qn('w:sectPr')]}
sect=deepcopy(body.sectPr)
for e in list(body): body.remove(e)
body.append(sect)

def p(text='',style=None): return doc.add_paragraph(text,style)
def h(text,level=1): return p(text,f'Heading {level}')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers))
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=str(v)
    if widths: t._custom_widths=widths
    return t
def field(paragraph,instruction):
    r=paragraph.add_run()
    for typ in ['begin','instr','separate','end']:
        el=OxmlElement('w:instrText' if typ=='instr' else 'w:fldChar')
        if typ=='instr': el.set(qn('xml:space'),'preserve'); el.text=instruction
        else: el.set(qn('w:fldCharType'),typ)
        r._r.append(el)

p('杆式夜间水体遥感反射率测量系统','Title')
p('室内分项实验综合报告','Subtitle')
p('光源特性与光谱仪性能验证','Subtitle')
p('测试探头  OCR-25F02100-1').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('编制依据  项目立项报告0826版v5').alignment=WD_ALIGN_PARAGRAPH.CENTER
p('本报告汇总现阶段室内分项实验及配套距离与光斑测试结果，服务于杆式夜间水体遥感反射率原理样机的光源选型、采集参数设置和质量控制。已测数据支持建立相对响应与候选参数，但尚不足以确认整机验收通过、最低可测辐亮度或可追溯绝对遥感反射率。')
h('主要结果',2)
p('光源方面，单灯冷启动记录显示短时波动较小，但约57 min内仍有缓慢下降；距离测试表明中心响应随距离增加而下降，横向光斑同时展宽，现有灯具需进一步核验实际工作区域的均匀性。两项结果的灯具身份及测试条件应分别管理，不能直接合并为同一光源的完整定型结论。')
p('探头方面，已测4000–45000 DN净光响应区间的积分时间比例偏差较小；信噪比测试中72 V / 64 ms为已测条件的优选候选设置。弱光结果表明延长积分时间有助于改善短波端表现，但最低已测设置仍未覆盖完整420–700 nm目标波段，绝对弱光边界尚需补测。')
h('目录',1)
field(p(), ' TOC \\o "1-2" \\h \\z \\u ')
h('1 项目目标与实验总体说明')
h('1.1 报告定位与技术主线',2)
p('立项报告以固定点、受控人工照明条件下的夜间表观遥感反射率测量为目标，测量链包括人工光源、同一探头的水体与参考板配对采集、暗场扣除、积分时间响应修正、照明差异修正及逐波长质量控制。本报告围绕这条测量链组织分项证据，先分析照明的时间和空间变化，再评价探头响应、信噪比与弱光可用性，最后归纳样机参数和待完成验证。')
p('报告以立项报告第8.2节规定的420–700 nm作为首期验收目标波段的对照范围。立项材料其他位置出现的420–720 nm及约415–739 nm作为扩展验证范围单列考虑；正式验收前应统一冻结口径。本报告不据扩展范围改变首期目标，也不将探头单侧的连续候选波段等同于水体与参考板共同有效波段。')
h('1.2 实验组成与数据关系',2)
p('表1-1  实验组成与数据关系','Caption')
table(['本报告章节','实验内容','数据关系与适用范围'],[
['第2章','单灯冷启动与输出稳定性','原实验A；一次冷启动记录。光源与原B、C、D不同。'],
['第3章','LED距离响应与横向光斑','独立距离测试；0.4 ms，1.507–5.013 m，中心及左右测点。'],
['第4章','积分时间响应比例','原实验B；6组长短积分配对，N模式数据。'],
['第5章','信噪比与连续候选波段','原实验C；5档电压、10组条件，逐组亮场与暗场。'],
['第6章','弱光边界与积分时间选择','原实验D；复用第5章数据，从弱信号角度分析，不计为独立重复验证。']])
p('距离实验与其他章节之间未建立可直接核验的灯号、供电和光学配置对应关系。因此，距离系数、冷启动等待建议及电压设置均只在各自试验条件内使用。后续将其组合为一套样机参数前，应以最终选定的同一灯具、探头和参考板完成联合复测。')
h('1.3 处理口径与物理量',2)
p('表1-2  数据处理口径与物理量','Caption')
table(['数据类别','处理口径','使用限制'],[
['冷启动实验','亮场减环境1固定遮挡背景；655–670 nm插值剔除定位笔干扰','DN·nm为相对积分响应，不能替代光度亮度或绝对辐亮度。'],
['距离与光斑实验','各测点亮场与同期遮阴光谱分别取中位数后相减','DN·channel为通道和；分析401.63–718.31 nm，不与DN·nm数值直接比较。'],
['积分响应实验','设备已扣暗的N模式导出值再减1000 DN偏置','不重复减独立暗场；实际积分时间按文件内部记录。'],
['信噪比与弱光实验','每组亮场减同积分时间平均暗谱','1000 DN偏置相减抵消；不再额外扣除，不用平滑光谱计算SNR。']])
p('遮挡光源直射光路得到的背景包含环境光与光路散射贡献，与遮蔽探头获得的仪器暗场不是同一量。各章保留相应处理方法，不将一种扣除方式替代另一种。未经辐射定标的DN响应不直接命名为绝对亮度、辐照度或辐亮度。')
h('1.4 判据层级与结果解释',2)
p('立项报告第5.7节为分项实验提出建议判据，第8.2节则面向完整原理样机。第5、6章的工作候选门槛为逐通道SNR≥20、码值不超过80%量程并满足采集期漂移要求；推荐候选进一步要求SNR≥50。整机验收中的水体与参考板共同SNR、6轮反射率重复性、自动测量成功率等需另行验证，不能由本报告的单侧探头统计代替。')
p('以下各章依次回答光源能否稳定使用、照明如何随空间位置变化、探头能否按积分时间换算，以及哪些通道在当前条件下可用。每章结论均限定在实际采集参数、记录长度和分析波段内。')

transitions={
'A':'光源的时间变化会直接影响水体与参考板先后测量的比值。本章先评价单灯冷启动后的响应变化，给出配对采集与参考监测的依据。',
'distance':'第2章讨论了光源随时间的变化；实际灯杆站还存在水位和安装几何变化。本章转向距离衰减与横向光斑，分析照明修正的依据以及可用照明区域。',
'B':'照明条件确定后，水体与参考板通常仍需采用不同积分时间。为判断两类响应能否换算到统一尺度，本章检验长、短积分时间配对的比例关系。',
'C':'积分时间比例接近理想关系，并不保证弱通道的测量质量。本章进一步联合考察信噪比、码值余量和采集期漂移，筛选各工况的连续候选波段。',
'D':'第5章确定了已测条件下的参数候选范围。本章复用同一批数据，重点回答光源减弱后哪些波段先失效、延长积分时间能够改善多少，以及最低可测亮度仍需哪些独立证据。'}
titles={'A':'2 单灯冷启动与输出稳定性','distance':'3 LED距离响应与横向光斑','B':'4 积分时间响应比例与工作码值','C':'5 信噪比与连续候选波段','D':'6 弱光性能与测量边界'}
chap={'A':2,'distance':3,'B':4,'C':5,'D':6}
figure_counts={}; table_counts={}
for key in ['A','distance','B','C','D']:
    n=chap[key]; h(titles[key]); p(transitions[key])
    fc=tc=0
    for e in sections[key]:
        if e.tag==qn('w:p'):
            pp=Paragraph(e,doc._body); txt=pp.text
            if txt.startswith('实验 ') or txt=='LED光源距离与光斑实验分析': continue
            if not txt.strip() and not e.xpath('.//w:drawing | .//m:oMath | .//w:pict'): continue
            if key=='distance' and txt.startswith('亮度够'):
                pp.text='现有窄角LED探照灯在近距离具有较高中心响应，但横向有效光斑较窄。在2–3 m工作距离下，现有测点不足以支持50–100 cm范围内的均匀照明；增加至4–5 m后光斑展宽，同时中心响应明显下降。灯具选择需同时评价目标区域的最低响应和中心信号强度。'
            if pp.style.name.startswith('Heading'):
                if key=='distance':
                    match=re.match(r'^(\d+)(?:\.(\d+))?\.?\s*(.*)',txt)
                    if match:
                        a,b,title=match.groups(); pp.text=f'3.{a}'+(f'.{b}' if b else '')+' '+title
                        pp.style=doc.styles['Heading 3' if b else 'Heading 2']
                else:
                    pp.text=re.sub(r'^'+key+r'\.',f'{n}.',txt)
            elif re.match(r'^图(?:A-|\d)',txt):
                fc+=1; pp.text=re.sub(r'^图(?:A-)?\d+\s*',f'图{n}-{fc}  ',txt); pp.style=doc.styles['Caption']
            else:
                replacements={
                'A.4.4':'2.4.4','图A-1':'图2-1','图A-3':'图2-3','B.4.2':'4.4.2',
                '实验B、C、D':'第4至6章','实验B':'第4章','实验C':'第5章','实验D':'第6章',
                '实验 B、C、D':'第4至6章'}
                for r in pp.runs:
                    for a,b in replacements.items():
                        if a in r.text: r.text=r.text.replace(a,b)
            # Use consistent notation without deleting original equation objects.
            if txt.startswith('式中DN_i'):
                pp.text='式中，DNᵢ（即Nᵢ）为第i帧亮场N模式输出，DNblk,j为环境1第j帧遮挡光谱，平均背景DN̄blk与N̄env1为同一量，净响应DNnet,i与Sᵢ为同一量。亮场和背景中的655–670 nm均已先插值处理。'
            if txt.startswith('（1）灯在约1 h'):
                pp.text='（1）本次灯具在约1 h内仍持续缓慢变暗。开灯约50 min后下降速率有所减小，可作为本灯后续试验的观察起点，但不能据此认定已达到热平衡或确定通用预热时间。每组测量前后应记录参考板响应，并复核稳定性后再采集。'
            if '由于六组规律一致，这属于可修正的系统偏差' in pp.text:
                pp.text=pp.text.replace('由于六组规律一致，这属于可修正的系统偏差，而非随机误差。','六组趋势一致提示存在系统性成分，但仪器响应与光源变化尚未分离；能否建立修正关系，仍需稳定光源和独立数据验证。')
            if '偏差最小主要因其位于系统性偏差趋势的过零区' in pp.text:
                pp.text=pp.text.replace('偏差最小主要因其位于系统性偏差趋势的过零区','偏差最小与其处于观测偏差趋势的过零区有关')
            body.insert(len(body)-1,e)
            if key=='distance' and txt.startswith('这些宽度均基于'):
                p('阈值总宽度与以中心对称截取的可用宽度不是同一指标。例如5.013 m的80%阈值总宽度约58.0 cm，是左右阈值交点距离之和；由于两侧分布不对称，仍可能出现±25 cm区域最低响应仅77.7%的情况。因此，不以58.0 cm总宽度直接判断居中的50 cm区域满足80%要求。')
        elif e.tag==qn('w:tbl'):
            tc+=1; p(f'表{n}-{tc}  '+tablemap[e],'Caption')
            body.insert(len(body)-1,e)
    figure_counts[n]=fc; table_counts[n]=tc

h('7 综合评价与样机应用建议')
h('7.1 与立项实验要求的对应关系',2)
p('现阶段成果可支撑样机的初始参数选择和补测设计。下表区分已经形成的分项证据与仍需补齐的验证，不以试验名称覆盖替代试验要求的实际完成。')
p('表7-1  分项证据与立项要求对照','Caption')
table(['立项要求','本报告证据','后续验证重点'],[
['单灯稳定性\n5.7.3','一次冷启动，约57 min记录；短时稳定、全程仍有漂移','至少3次独立冷启动、10次启停；记录绝对亮度并确定Tstable分布。'],
['距离与径向光斑\n5.7.4','5个距离，中心与左右测点；幂律指数1.8706','非单调距离序列与3轮重复；上下及必要斜向测点、边缘加密；验证查找表。'],
['光源光谱特性\n5.7.5','部分时间、距离和位置的相对谱形比较','按规定时点补齐光谱角、NRMSE和峰位；655–670 nm插值区单独标记。'],
['多灯长期一致性\n5.7.6','尚无对应的完整多灯长期验证结果','同台架多灯连续7 d测试，主备灯排序与替换系数验证。'],
['积分响应\n5.7.7','6组两档积分配对；中间码值比例偏差较小','稳定光源下多档积分回归；独立验证残差、适用波段和饱和边界。'],
['信噪比\n5.7.8','10组条件，每组30或50帧；提供逐通道联合筛选','补齐实测目标亮度、100帧序列与独立重复；核验正式目标波段。'],
['波长与绝对辐射定标\n5.7.9至5.7.10','尚无本组实验的独立定标验证结果','带证书谱线验证；按测量路线需求完成辐射定标或限定相对量输出。'],
['弱光性能\n5.7.11','5档电压、最长128 ms；最低档仍未覆盖420–700 nm','实测亮度不少于6档，覆盖工程最低值以下；确定Lmin、积分上限和平均帧数。']])
p('此外，立项报告要求的同一工况至少3轮独立重复、温度与供电记录、参考板信息等，应在后续统一记录表中补齐。现有报告未给出完整证据的项目，不按已完成验收处理。')
h('7.2 初始参数及其使用条件',2)
p('表7-2  初始参数及适用条件','Caption')
table(['应用环节','已测依据与初始建议','使用条件'],[
['照明几何','3.995 m可作为现有灯具下兼顾信号与横向宽度的候选距离','不是推荐安装高度；需按真实视场、入射角、水位及目标区域联合验证。'],
['距离修正','1.507–5.013 m范围可用实测响应及幂律趋势作初始参考','不得外推；逐波长照明差异Cd仍需水面与板位置实测或验证。'],
['净光工作码值','已测4000–45000 DN区间比例偏差较小；25000–35000 DN可作试验参考','不等同于普适线性区；另行检查亮场量程占用及所有目标通道。'],
['常规信号采集','72 V / 64 ms为第5章已测条件的优选候选','仅限该灯具与供电配置；现场重新预采集，不能用于其他灯具供电。'],
['较弱信号采集','60 V条件下128 ms优于64 ms','128 ms不是已验证积分上限；该条件连续工作波段仍从474.6 nm起。'],
['波段输出','按净光正值、SNR、量程及漂移联合筛选，保留连续区间和原因码','水体和参考板分别筛选后取交集；冷启动实验中655–670 nm插值区单独标记。']])
h('7.3 从分项实验到水体配对测量',2)
p('第一步，固定最终灯具、探头、参考板及收发几何，为每套配置建立可追溯编号。对最终光源重新完成预热与光斑验证，明确实际水面和参考板所在位置，避免将不同灯具、不同测量场景的参数直接拼接。')
p('第二步，以现场水体与参考板预采集结果选择各自积分时间，在同积分时间与相近温度下采集暗场；背景测量另行记录。积分切换后等待稳态，保留原始帧、时间戳、温度、灯号、电压及实际距离，防止重复扣暗或混用码值定义。')
p('第三步，在已验证积分响应范围内归一化水体与参考板信号，使用参考板逐波长反射率和经实测验证的照明差异系数计算夜间表观遥感反射率。参考板与水面的有效通道取交集，并检查配对时间、距离变化、饱和和镜面反射。')
p('第四步，在固定水体上完成至少6轮配对重复性验证，再开展室外、日夜衔接与同步水质实验。现有单侧探头SNR和中心重复差分别描述不同量，不用于替代Rrs的轮次CV、日夜谱形相关或水质反演性能。')
h('8 综合结论与后续工作')
p('现阶段实验已形成光源时间变化、距离与横向光斑响应、积分时间比例关系、信噪比与弱光候选范围五方面证据，可用于原理样机下一阶段的配置筛选。现有窄角光源的主要工程约束是中心信号与均匀照明范围之间的取舍；探头参数则需同时满足弱通道SNR和强通道量程余量。')
p('在本次积分响应数据中，4000–45000 DN范围的入选数据点最大绝对比例偏差为1.738%；在信噪比数据中，72 V / 64 ms的最长连续工作候选波段为409.0–878.1 nm，覆盖420–700 nm对照范围，但其推荐候选波段从447.2 nm起。这些结果限定于相应测试光源和测量序列，不能据此宣称真实水体条件下已经达到整机目标。')
p('弱光下，60 V / 128 ms的最长连续工作候选波段为474.6–878.1 nm，仍不足以覆盖420–700 nm。由于缺少实测辐亮度和更低亮度档位，最低可测辐亮度、积分时间上限及最优平均帧数尚未确定；669 DN的外推值仅用于探索，不进入样机性能指标。')
p('下一阶段宜优先完成最终灯具的稳定性与二维光斑复测、稳定光源下的多档积分响应、实际亮度分级的100帧信噪比与弱光验证，并补齐波长准确度及必要的辐射定标。在此基础上开展水体与参考板配对闭环，再依据立项报告验收指标评价样机，输出名称保持为夜间表观遥感反射率。')

# Clean prior review residue and manual page controls; preserve all figure and equation objects.
for tag in ['del','commentRangeStart','commentRangeEnd','commentReference','lastRenderedPageBreak']:
    for e in doc.element.xpath('//w:'+tag): e.getparent().remove(e)
for e in doc.element.xpath('//w:ins'):
    par=e.getparent(); idx=par.index(e)
    for child in list(e): par.insert(idx,child); idx+=1
    par.remove(e)
for e in doc.element.xpath('//w:br[@w:type="page"]'): e.getparent().remove(e)
for e in doc.element.xpath('//w:pPr/w:numPr'): e.getparent().remove(e)

def font(run,size=12,bold=None,east='宋体'):
    run.font.name='Times New Roman'; run.font.size=Pt(size); run.font.color.rgb=RGBColor(0,0,0)
    run.italic=False
    if bold is not None: run.bold=bold
    rf=run._r.get_or_add_rPr().get_or_add_rFonts()
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+k),None)
    rf.set(qn('w:eastAsia'),east)
for name,size,east in [('Normal',12,'宋体'),('Title',22,'黑体'),('Subtitle',16,'黑体'),('Heading 1',16,'黑体'),('Heading 2',14,'黑体'),('Heading 3',12,'黑体'),('Caption',10.5,'宋体')]:
    st=doc.styles[name]; st.font.name='Times New Roman'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0); st.font.italic=False
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)
    st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(6)
    if name.startswith('Heading'): st.font.bold=True
    for e in st._element.xpath('.//w:pBdr'): e.getparent().remove(e)
for sec in doc.sections:
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.3); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.4); sec.right_margin=Cm(2.4)
    sec.header_distance=Cm(1); sec.footer_distance=Cm(1)
    for pp in sec.header.paragraphs: pp.clear()
    for pp in sec.footer.paragraphs: pp.clear()
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field(foot,' PAGE ')
    for r in foot.runs: font(r,9)

list_count=0
for pp in doc.paragraphs:
    if pp.style.name.startswith('Heading'): list_count=0
    if pp.style.name.startswith('List Number'):
        list_count+=1; pp.text=f'（{list_count}）'+pp.text; pp.style=doc.styles['Normal']
    elif pp.style.name.startswith('List Bullet'):
        pp.style=doc.styles['Normal']
    pf=pp.paragraph_format
    pf.page_break_before=False; pf.keep_with_next=False; pf.keep_together=False; pf.widow_control=True
    pf.left_indent=Pt(0); pf.right_indent=Pt(0); pf.first_line_indent=Pt(24)
    pf.line_spacing=1.35; pf.space_before=Pt(0); pf.space_after=Pt(5)
    pp.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for e in pp._p.xpath('./w:pPr/w:pBdr'): e.getparent().remove(e)
    name=pp.style.name; txt=pp.text
    size=12; east='宋体'; bold=False
    if name=='Title':
        size=22; east='黑体'; bold=True; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.space_before=Pt(36); pf.space_after=Pt(18); pf.first_line_indent=Pt(0)
    elif name=='Subtitle':
        size=16; east='黑体'; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.space_after=Pt(14); pf.first_line_indent=Pt(0)
    elif name.startswith('Heading'):
        lev=int(name[-1]); size={1:16,2:14,3:12}[lev]; east='黑体'; bold=True
        pf.first_line_indent=Pt(0); pf.keep_with_next=True; pf.keep_together=True; pf.space_before=Pt(12); pf.space_after=Pt(7)
        if lev==1: pf.page_break_before=True
    elif name=='Caption':
        size=10.5; pf.first_line_indent=Pt(0); pf.line_spacing=1.15; pf.space_after=Pt(6); pf.keep_together=True
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.keep_with_next=txt.startswith('表')
    elif pp._p.xpath('.//w:drawing | .//w:pict'):
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Pt(0); pf.line_spacing=1; pf.keep_with_next=True
    elif pp._p.xpath('.//m:oMath') or (len(txt)<155 and any(x in txt for x in [' = ','SNR(', 'CV(', 'γ=', 'δ ='])):
        pf.first_line_indent=Pt(0); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.keep_together=True; size=11
    if txt in ['目录','主要结果']:
        pp.style=doc.styles['Normal']
        ol=OxmlElement('w:outlineLvl'); ol.set(qn('w:val'),'9'); pp._p.get_or_add_pPr().append(ol)
    if txt.endswith(('：',':')) and len(txt)<160: pf.keep_with_next=True
    if txt.startswith(('测试探头','编制依据')): pf.first_line_indent=Pt(0); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in pp.runs: font(r,size,bold,east)
    if not txt.strip() and pp._p.xpath('.//w:instrText'):
        pf.first_line_indent=Pt(0)

# Table widths, repeat headers, and non-splitting rows.
usable=16.2
for t in doc.tables:
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    cols=len(t.columns)
    headers=[c.text for c in t.rows[0].cells]
    if cols==2: widths=[4.0,12.2]
    elif cols==3: widths=[3.0,6.3,6.9] if any('后续' in x or '适用' in x or '处理口径' in x or '初始建议' in x for x in headers) else [5.4]*3
    elif cols==5 and '最长连续' in '|'.join(headers): widths=[3,5,1.6,5,1.6]
    elif cols==8: widths=[3.3]+[12.9/7]*7
    else: widths=[usable/cols]*cols
    for grid,w in zip(t._tbl.tblGrid.gridCol_lst,widths): grid.set(qn('w:w'),str(round(w/2.54*1440)))
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        for el in list(trpr):
            if el.tag in [qn('w:trHeight'),qn('w:cantSplit'),qn('w:tblHeader')]: trpr.remove(el)
        trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            for e in list(pr):
                if e.tag in [qn('w:shd'),qn('w:tcMar'),qn('w:tcBorders')]: pr.remove(e)
            mar=OxmlElement('w:tcMar')
            for side,val in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),val); el.set(qn('w:type'),'dxa'); mar.append(el)
            pr.append(mar)
            if ri==0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'E8ECF0'); pr.append(shd)
            for pp in c.paragraphs:
                pf=pp.paragraph_format; pf.first_line_indent=Pt(0); pf.left_indent=Pt(0); pf.right_indent=Pt(0); pf.line_spacing=1.1
                pf.space_before=Pt(0); pf.space_after=Pt(2); pf.keep_with_next=(ri==0); pf.keep_together=False
                pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if cols<=3 else WD_ALIGN_PARAGRAPH.CENTER
                for r in pp.runs: font(r,9 if cols>=6 else 10,ri==0)
    pr=t._tbl.tblPr
    for e in list(pr):
        if e.tag in [qn('w:tblBorders'),qn('w:tblInd'),qn('w:tblCellSpacing')]: pr.remove(e)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'B8BFC6'); borders.append(el)
    pr.append(borders)
for shape in doc.inline_shapes:
    if shape.width>Cm(16.0):
        ratio=Cm(16.0)/shape.width; shape.width=Cm(16.0); shape.height=int(shape.height*ratio)
    if shape.height>Cm(13.4):
        ratio=Cm(13.4)/shape.height; shape.height=Cm(13.4); shape.width=int(shape.width*ratio)
doc.core_properties.title='杆式夜间水体遥感反射率测量系统室内分项实验综合报告'
doc.core_properties.subject='对照立项报告整理的光源特性与光谱仪性能验证'
doc.core_properties.comments=''
out=BASE.parent/'室内分项实验综合报告_按立项报告整理.docx'
doc.save(out)
print(json.dumps({'output':str(out),'tables':len(doc.tables),'images':len(doc.inline_shapes),'equations':len(doc.element.xpath('//m:oMath')),'figures_by_chapter':figure_counts},ensure_ascii=False))
