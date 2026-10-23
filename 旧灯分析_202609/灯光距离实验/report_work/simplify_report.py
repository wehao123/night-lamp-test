from pathlib import Path
from docx import Document
from copy import deepcopy

base=Path(__file__).resolve().parent.parent
doc=Document(base/'室内分项实验综合报告_420–720nm修订版.docx')
def replace(p,text):
    prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if prop is not None: r._r.insert(0,prop)
def remove(e): e.getparent().remove(e)
for t in list(doc.tables):
    if t.cell(0,0).text in ['本报告章节','立项要求']: remove(t._element)
in_removed_section=False
for p in list(doc.paragraphs):
    text=p.text
    if p.style.name.lower().startswith('toc'): continue
    if text=='7.1 与立项实验要求的对应关系': in_removed_section=True
    if text=='7.2 初始参数及其使用条件': in_removed_section=False
    if in_removed_section:
        remove(p._p); continue
    if text.startswith(('编制依据','表1-1  实验组成')):
        remove(p._p); continue
    edits={
    '1 项目目标与实验总体说明':'1 实验概述',
    '1.1 报告定位与技术主线':'1.1 实验目的与目标波段',
    '1.2 实验组成与数据关系':'1.2 实验内容',
    '1.3 处理口径与物理量':'1.3 数据处理与物理量',
    '1.4 判据层级与结果解释':'1.4 评价指标',
    '表1-2  数据处理口径与物理量':'表1-1  数据处理方法与物理量',
    '7.2 初始参数及其使用条件':'7.1 初始参数及其使用条件',
    '表7-2  初始参数及适用条件':'表7-1  初始参数及适用条件',
    '7.3 从分项实验到水体配对测量':'7.2 水体配对测量建议',
    }
    if text in edits: replace(p,edits[text]); continue
    if text.startswith('本报告汇总现阶段室内分项实验'):
        replace(p,'本报告汇总LED光源与光谱探头的分项实验结果，分析光源输出稳定性、距离衰减与光斑分布，以及探头的积分时间响应、信噪比和弱光性能，为夜间水体光谱测量选择照明条件、采集参数与有效波段。')
    elif text.startswith('光源方面，单灯冷启动记录'):
        replace(p,'光源方面，单灯冷启动记录显示短时波动较小，但约57 min内仍有缓慢下降；距离测试表明中心响应随距离增加而下降，横向光斑同时展宽。光源使用时需兼顾输出稳定性、目标区域的照明均匀性和有效信号强度。')
    elif text.startswith('立项报告以固定点'):
        replace(p,'夜间水体光谱测量使用LED照明，并通过同一探头先后测量水体和参考板。光源随时间和位置的变化、探头对积分时间的响应，以及弱信号中的噪声，都会影响两次测量的可比性。本报告分别评价这些因素，确定现有设备的可用条件及主要限制。')
    elif text.startswith('本项目正式目标波段统一'):
        replace(p,'本项目目标波段为420–720 nm。现有冷启动实验的积分统计范围为420–700 nm，其结果仍按实际计算范围报告，700–720 nm的稳定性需另行评价。其他实验按各章注明的分析波段处理，并结合420–720 nm目标范围判断测量可用性。')
    elif text.startswith('距离实验与其他章节之间未建立'):
        replace(p,'光源实验包括两部分：冷启动与输出稳定性实验，观察灯具点亮后的短时波动和持续漂移；距离与光斑实验，测量不同灯距下的中心响应和横向分布，判断照明衰减及均匀区域的大小。')
        q=deepcopy(p._p); p._p.addnext(q)
        from docx.text.paragraph import Paragraph
        replace(Paragraph(q,p._parent),'探头实验包括三部分：积分时间响应实验，检验长、短积分时间的信号能否按比例换算；信噪比实验，结合帧间波动、码值余量和漂移筛选可用通道；弱光分析，考察信号减弱后延长积分时间的改善效果及测量边界。各章分别说明实验条件、处理方法、结果和结论。')
    elif text.startswith('立项报告第5.7节为分项实验'):
        replace(p,'光源稳定性以滑动变异系数、变化率和长时间响应变化评价；距离与光斑以中心响应、归一化横向分布和阈值宽度评价；积分时间响应以长、短积分的比例偏差评价。信噪比与弱光分析按逐通道SNR、码值余量和采集期漂移联合筛选：工作候选要求SNR≥20，推荐候选要求SNR≥50，并同时满足各章给出的其余条件。')
    elif text.startswith('下一阶段宜优先完成最终灯具'):
        replace(p,'下一阶段宜优先完成最终灯具的稳定性与二维光斑复测、稳定光源下的多档积分响应、实际亮度分级的100帧信噪比与弱光验证，并补齐波长准确度及必要的辐射定标。在此基础上开展水体与参考板配对测量，评价420–720 nm内的共同有效波段与重复性，输出名称保持为夜间表观遥感反射率。')
doc.core_properties.subject='LED光源与光谱探头分项实验结果'
out=base/'室内分项实验综合报告_420–720nm修订版_v2.docx'
doc.save(out)
print(out)
